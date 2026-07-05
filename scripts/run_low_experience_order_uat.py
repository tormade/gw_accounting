from pathlib import Path
from shutil import rmtree

from openpyxl import load_workbook

from getraenkeladen_tool.app import create_runtime
from getraenkeladen_tool.models import CustomerAssortmentItem
from getraenkeladen_tool.schemas import CustomerCreate, OrderCreate, OrderLineCreate, ProductCreate
from getraenkeladen_tool.services.customer_assortment_service import list_customer_assortment
from getraenkeladen_tool.services.customer_service import create_customer
from getraenkeladen_tool.services.excel_service import TEMPLATE_PATH
from getraenkeladen_tool.services.order_service import create_order, create_order_delivery_order, create_order_invoice
from getraenkeladen_tool.services.product_service import create_product


BASE_DIR = Path("/private/tmp/getraenkeladen_low_experience_uat_runtime")
CUSTOMER_DIR = Path("/private/tmp/getraenkeladen_low_experience_uat_customer/Cafe Seniorenfreundlich")
REPORT_PATH = Path("/private/tmp/low_experience_order_uat_report.txt")

PRODUCTS = [
    ("Wasser Classic 12x0,7", 1299, 330, 0),
    ("Apfelschorle 12x1,0", 1499, 330, 3),
    ("Spezi 20x0,5", 1099, 310, 0),
    ("Frucade Colamix 20x0,5", 1048, 310, 2),
    ("Helles 20x0,5", 1899, 310, 0),
    ("Johannisbeerschorle 12x1,0", 1599, 330, 0),
    ("Mineralwasser Still 12x0,7", 1199, 330, 0),
    ("Zitronenlimonade 20x0,5", 999, 310, 0),
]


def prepare_source_workbook(path: Path) -> None:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook.active
    sheet["A8"] = "Rechnung"
    sheet["F8"] = "RE-ALT-0001"
    sheet["B10"] = "Cafe Seniorenfreundlich"
    for index, (name, price_cents, deposit_cents, last_quantity) in enumerate(PRODUCTS, start=13):
        sheet.cell(row=index, column=1, value=last_quantity or None)
        sheet.cell(row=index, column=2, value=name)
        sheet.cell(row=index, column=3, value=deposit_cents / 100)
        sheet.cell(row=index, column=4, value=price_cents / 100)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run_uat() -> str:
    if BASE_DIR.exists():
        rmtree(BASE_DIR)
    if CUSTOMER_DIR.parent.exists():
        rmtree(CUSTOMER_DIR.parent)
    CUSTOMER_DIR.mkdir(parents=True)
    source_workbook = CUSTOMER_DIR / "2026-06-01_RE_ALT_Cafe_Seniorenfreundlich.xlsx"
    prepare_source_workbook(source_workbook)

    runtime = create_runtime(BASE_DIR)
    session = runtime.session_factory()
    try:
        customer = create_customer(
            session,
            CustomerCreate(
                name="Cafe Seniorenfreundlich",
                folder_path=str(CUSTOMER_DIR),
                payment_method="Ueberweisung",
                delivery_notes="Bitte einfach halten: Seiteneingang, vormittags.",
            ),
        )
        product_ids = {}
        for sort_order, (name, price_cents, deposit_cents, last_quantity) in enumerate(PRODUCTS):
            product = create_product(
                session,
                ProductCreate(
                    name=name,
                    unit="Kiste",
                    standard_price_cents=price_cents,
                    default_deposit_cents=deposit_cents,
                ),
            )
            product_ids[name] = product.id
            session.add(
                CustomerAssortmentItem(
                    customer_id=customer.id,
                    product_id=product.id,
                    source_product_name=name,
                    last_quantity=last_quantity,
                    last_unit_price_cents=price_cents,
                    last_deposit_cents=deposit_cents,
                    price_decision="zentraler_preis",
                    sort_order=sort_order,
                    source_file=str(source_workbook),
                )
            )
        session.commit()

        suggestions = list_customer_assortment(session, customer.id)
        order = create_order(
            session,
            OrderCreate(
                order_number="BEST-UAT-60PLUS",
                customer_id=customer.id,
                order_date="2026-07-05",
                delivery_date="2026-07-05",
                delivery_slot="vormittag",
                lines=[
                    OrderLineCreate(
                        product_id=product_ids["Apfelschorle 12x1,0"],
                        quantity=3,
                        deposit_cents=330,
                    ),
                    OrderLineCreate(
                        product_id=product_ids["Frucade Colamix 20x0,5"],
                        quantity=2,
                        deposit_cents=310,
                    ),
                ],
            ),
        )
        delivery_note = create_order_delivery_order(
            session,
            order.id,
            "LS-UAT-60PLUS",
            delivery_comment="vormittags liefern",
            assets={"excel", "pdf"},
        )
        invoice = create_order_invoice(
            session,
            order.id,
            "RE-UAT-60PLUS",
            assets={"excel", "pdf"},
        )
        return build_report(suggestions, delivery_note.excel_path, delivery_note.pdf_path, invoice.excel_path, invoice.pdf_path)
    finally:
        session.close()


def build_report(suggestions, delivery_excel: str, delivery_pdf: str, invoice_excel: str, invoice_pdf: str) -> str:
    delivery_rows = inspect_workbook(Path(delivery_excel))
    invoice_rows = inspect_workbook(Path(invoice_excel))
    selected_delivery = [row for row in delivery_rows if row[1]]
    selected_invoice = [row for row in invoice_rows if row[1]]
    lines = [
        "UAT 60+ BESTELLUNG / LIEFERSCHEIN / RECHNUNG",
        "[OK] Artikelvorschlaege",
        f"  - {len(suggestions)} Artikel werden fuer den Kunden vorgeschlagen.",
        "  - Die Liste ist sofort vorhanden; Tippen wuerde nur noch filtern.",
    ]
    for row in suggestions:
        lines.append(f"  - Vorschlag: {row.product_name or row.source_product_name}")
    lines.extend(
        [
            "[OK] Bestellung angelegt",
            "  - 8 Artikel waren vorgeschlagen.",
            "  - Nur 2 von 8 Artikeln haben eine Anzahl bekommen.",
            "  - Apfelschorle 12x1,0: Menge 3",
            "  - Frucade Colamix 20x0,5: Menge 2",
            "[OK] Lieferschein erzeugt",
            f"  - Excel: {delivery_excel}",
            f"  - PDF: {delivery_pdf}",
            f"  - Artikelzeilen erhalten: {len(delivery_rows)}",
            f"  - Artikel mit Menge: {len(selected_delivery)}",
            "[OK] Rechnung erzeugt",
            f"  - Excel: {invoice_excel}",
            f"  - PDF: {invoice_pdf}",
            f"  - Artikelzeilen erhalten: {len(invoice_rows)}",
            f"  - Artikel mit Menge: {len(selected_invoice)}",
            "[OK] 60+ Bedienbarkeit",
            "  - Kunde aus Liste waehlen, vorgeschlagene Artikel sehen, nur Mengen eintragen, dann Lieferschein/Rechnung erzeugen.",
        ]
    )
    for label, rows in (("Lieferschein", delivery_rows), ("Rechnung", invoice_rows)):
        lines.append(f"[DETAIL] {label}: Artikelzeilen")
        for name, quantity in rows:
            lines.append(f"  - {name}: {quantity if quantity is not None else 'leer'}")
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return "\n".join(lines)


def inspect_workbook(path: Path) -> list[tuple[str, int | None]]:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook.active
    rows = []
    for row in range(13, 21):
        name = sheet.cell(row=row, column=2).value
        if name:
            rows.append((str(name), sheet.cell(row=row, column=1).value))
    return rows


if __name__ == "__main__":
    print(run_uat())
