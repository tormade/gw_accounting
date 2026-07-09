from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEMO_DIR = ROOT / "Demo_Pruefliste"
INPUT_DIR = DEMO_DIR / "Input"
FOLDER_DIR = DEMO_DIR / "Kundenordner"


ARTICLES = [
    ["Prudukt Bezeichnung", "Fl.", "Menge", "Liefer Preis", "Pfand", "Demo-Hinweis"],
    ["Paulaner Spezi 20x0,5", 20, "0,5", 12.50, 3.10, "In Kundenordner bewusst anderer Preis"],
    ["Adelholzener Classic 12x0,7", 12, "0,7", 8.50, 3.30, "Soll ohne Pruefpunkt passen"],
    ["Frucade Colamix 20x0,5", 20, "0,5", 10.48, 3.10, "In Kundenordner bewusst anderer Preis"],
    ["Orangina 24x0,25", 24, "0,25", 18.90, 3.42, "Soll ohne Pruefpunkt passen"],
]


CUSTOMERS = [
    [
        "Name",
        "Straße",
        "Hausnummer",
        "PLZ",
        "Ort",
        "Ansprechpartner",
        "e-mail (Kontakt)",
        "e-mail (Re Versand)",
        "Festnetz",
        "nächster Kontakt",
        "Bemerkungen",
        "ABO ",
        "Vormittag",
        "Nachmittag",
    ],
    [
        "Demo Wirtshaus",
        "Hauptstrasse",
        "1",
        "85354",
        "Freising",
        "Frau Maier",
        "kontakt@demo-wirtshaus.de",
        "rechnung@demo-wirtshaus.de",
        "08161 111111",
        date(2026, 7, 1),
        "Lieferung ueber Hintereingang",
        "14-taegig",
        "x",
        "",
    ],
    [
        "Demo Vereinsheim",
        "Sportplatzweg",
        "7",
        "85406",
        "Zolling",
        "Herr Huber",
        "kontakt@demo-verein.de",
        "rechnung@demo-verein.de",
        "08167 222222",
        date(2026, 7, 3),
        "Vorher anrufen",
        "monatlich",
        "",
        "x",
    ],
]


def main() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    (FOLDER_DIR / "Demo Wirtshaus").mkdir(parents=True, exist_ok=True)
    (FOLDER_DIR / "Demo Vereinsheim").mkdir(parents=True, exist_ok=True)

    _write_table_workbook(INPUT_DIR / "Artikel Liste Preise.xlsx", ARTICLES)
    _write_table_workbook(INPUT_DIR / "Lieferkunden Liste.xlsx", CUSTOMERS)

    _write_customer_workbook(
        FOLDER_DIR / "Demo Wirtshaus" / "2026-06-24_RE_DEMO-1001_Demo_Wirtshaus.xlsx",
        customer_name="Demo Wirtshaus",
        address_lines=["Falsche Strasse 99", "85354 Freising"],
        email="alt-rechnung@demo-wirtshaus.de",
        phone="08161 999999",
        document_number="DEMO-1001",
        document_date=date(2026, 6, 24),
        comment="bis 13 Uhr und ab 15 Uhr",
        footer="Rechnungsbetrag wird per SEPA-Basis Lastschriftmandat eingezogen.",
        lines=[
            ("Paulaner Spezi 20x0,5", 2, 3.10, 11.90),
            ("Mystery Cola 20x0,5", 1, 3.10, 9.99),
            ("Adelholzener Classic 12x0,7", 3, 3.30, 8.50),
        ],
        deposit_returns=[(1, 3.10)],
        delivery_fee=3.90,
    )
    _write_customer_workbook(
        FOLDER_DIR / "Demo Wirtshaus" / "2026-05-10_RE_DEMO-0999_Demo_Wirtshaus_ALT.xlsx",
        customer_name="Demo Wirtshaus",
        address_lines=["Hauptstrasse 1", "85354 Freising"],
        email="rechnung@demo-wirtshaus.de",
        phone="08161 111111",
        document_number="DEMO-0999",
        document_date=date(2026, 5, 10),
        comment="alte Datei, soll beim Ordnerimport uebersprungen werden",
        footer="Rechnungsbetrag wird per SEPA-Basis Lastschriftmandat eingezogen.",
        lines=[("Adelholzener Classic 12x0,7", 1, 3.30, 8.50)],
        deposit_returns=[],
        delivery_fee=0,
    )
    _write_customer_workbook(
        FOLDER_DIR / "Demo Vereinsheim" / "2026-06-23_LS_DEMO-2001_Demo_Vereinsheim.xlsx",
        customer_name="Demo Vereinsheim",
        address_lines=["Sportplatzweg 77", "85406 Zolling"],
        email="vorstand@demo-verein.de",
        phone="08167 333333",
        document_number="DEMO-2001",
        document_date=date(2026, 6, 23),
        comment="nur Freitag Nachmittag",
        footer="Die Ware bleibt bis zur vollstaendigen Bezahlung Eigentum von Getraenke Winklmeier.",
        lines=[
            ("Frucade Colamix 20x0,5", 4, 3.10, 10.99),
            ("Club Mate 20x0,5", 2, 3.10, 14.50),
            ("Orangina 24x0,25", 1, 3.42, 18.90),
        ],
        deposit_returns=[(2, 3.10), (1, 3.42)],
        delivery_fee=0,
    )
    lock_file = FOLDER_DIR / "Demo Vereinsheim" / "~$2026-06-23_LS_DEMO-2001_Demo_Vereinsheim.xlsx"
    lock_file.write_text("Demo-Sperrdatei: muss beim Import uebersprungen werden.\n", encoding="utf-8")
    _write_readme()

    print(f"Demo-Prueflisten-Dateien erstellt: {DEMO_DIR}")


def _write_table_workbook(path: Path, rows: list[list]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tabelle1"
    for row in rows:
        sheet.append(row)
    _format_sheet(sheet)
    workbook.save(path)


def _write_customer_workbook(
    path: Path,
    *,
    customer_name: str,
    address_lines: list[str],
    email: str,
    phone: str,
    document_number: str,
    document_date: date,
    comment: str,
    footer: str,
    lines: list[tuple[str, int, float, float]],
    deposit_returns: list[tuple[int, float]],
    delivery_fee: float,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tabelle1"
    sheet["A1"] = "Firma"
    sheet["B1"] = comment
    sheet["A2"] = customer_name
    sheet["A3"] = address_lines[0]
    sheet["A4"] = address_lines[1]
    sheet["D2"] = document_date
    sheet["D2"].number_format = "yyyy-mm-dd"
    sheet["A7"] = email
    sheet["A8"] = phone
    sheet["A9"] = "Re. Nr."
    sheet["B9"] = document_number

    sheet["A12"] = "Menge"
    sheet["B12"] = "Artikel"
    sheet["C12"] = "Pfand"
    sheet["D12"] = "Lieferpreis"
    sheet["E12"] = "Zeilensumme"
    for offset, (name, quantity, deposit, price) in enumerate(lines, start=13):
        sheet.cell(row=offset, column=1, value=quantity)
        sheet.cell(row=offset, column=2, value=name)
        sheet.cell(row=offset, column=3, value=deposit)
        sheet.cell(row=offset, column=4, value=price)
        sheet.cell(row=offset, column=5, value=round(quantity * (deposit + price), 2))

    delivery_total = round(sum(quantity * (deposit + price) for _, quantity, deposit, price in lines), 2)
    return_total = -round(sum(quantity * deposit for quantity, deposit in deposit_returns), 2)
    gross_total = round(delivery_total + return_total + delivery_fee, 2)
    net_total = round(gross_total / 1.19, 2)
    tax_total = round(gross_total - net_total, 2)

    sheet["A31"] = "Lieferpauschale"
    sheet["E31"] = delivery_fee
    sheet["A32"] = "Lieferwert"
    sheet["B32"] = sum(quantity for _, quantity, _, _ in lines)
    sheet["E32"] = delivery_total

    sheet["A33"] = "Pfand R\u00fcckgabe"
    sheet["E33"] = return_total
    for offset, (quantity, deposit) in enumerate(deposit_returns, start=34):
        sheet.cell(row=offset, column=1, value=quantity)
        sheet.cell(row=offset, column=2, value="Pfandrueckgabe")
        sheet.cell(row=offset, column=3, value=deposit)
        sheet.cell(row=offset, column=5, value=-round(quantity * deposit, 2))
    sheet["A40"] = "Netto"
    sheet["E40"] = net_total
    sheet["A41"] = "Mehrwertsteuer"
    sheet["E41"] = tax_total
    sheet["A42"] = "Brutto"
    sheet["E42"] = gross_total
    sheet["A44"] = footer

    for row in range(13, 43):
        for col in (3, 4, 5):
            sheet.cell(row=row, column=col).number_format = '#,##0.00 "EUR"'
    _format_sheet(sheet)
    workbook.save(path)


def _format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in sheet[1]:
        if cell.value is not None:
            cell.font = Font(bold=True)
            cell.fill = header_fill
    for cell in sheet[12]:
        if cell.value is not None:
            cell.font = Font(bold=True)
            cell.fill = header_fill
    for column in range(1, 7):
        sheet.column_dimensions[get_column_letter(column)].width = [14, 36, 12, 14, 16, 18][column - 1]
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"


def _write_readme() -> None:
    readme = DEMO_DIR / "README.md"
    readme.write_text(
        "\n".join(
            [
                "# Demo-Pruefliste",
                "",
                "Diese Dateien sind bewusst fehlerhaft bzw. widerspruechlich, damit die Pruefliste gefuellt wird.",
                "",
                "## Nutzung in der App",
                "",
                "1. In `Einstellungen` den Input-Ordner auf `Demo_Pruefliste/Input` setzen.",
                "2. Stammdaten importieren.",
                "3. Den Kundenordner-Import auf `Demo_Pruefliste/Kundenordner` setzen.",
                "4. Kundenordner importieren.",
                "5. `Pruefpunkte ansehen` oeffnen.",
                "",
                "## Erwartete Pruefpunkte",
                "",
                "- Demo Wirtshaus: Adresse, Telefon und E-Mail weichen zwischen Lieferkundenliste und Ordnerdatei ab.",
                "- Demo Wirtshaus: `Paulaner Spezi 20x0,5` hat einen abweichenden Lieferpreis.",
                "- Demo Wirtshaus: `Mystery Cola 20x0,5` ist nicht im Artikelstamm vorhanden.",
                "- Demo Vereinsheim: Adresse, Telefon und E-Mail weichen ab.",
                "- Demo Vereinsheim: `Frucade Colamix 20x0,5` hat einen abweichenden Lieferpreis.",
                "- Demo Vereinsheim: `Club Mate 20x0,5` ist nicht im Artikelstamm vorhanden.",
                "- Eine alte Datei und eine Excel-Sperrdatei werden beim Ordnerimport uebersprungen.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
