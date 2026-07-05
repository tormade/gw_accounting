from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from ..kern.regeln.beleg import BelegParameter, PfandRueckgabe, Position, berechne_beleg
from .file_service import ensure_parent_folder


TEMPLATE_PATH = Path(__file__).resolve().parents[3] / "templates" / "vorlage_liefern_bar.xlsx"
FIRST_ITEM_ROW = 13
MAX_ITEM_ROW = 30
FIRST_RETURN_ROW = 34
MAX_RETURN_ROW = 39
DELIVERY_FEE_CENTS = 390
DELIVERY_NOTE_FOOTER = "Die Ware bleibt bis zur vollstaendigen Bezahlung Eigentum von Getraenke Winklmeier."
THANK_YOU_TEXT = "Vielen Dank fuer Ihren Einkauf"


def build_invoice_workbook(
    output_path: Path,
    customer_name: str,
    document_number: str,
    line_items: list[dict],
    document_date: str | None = None,
    deposit_returns: list[dict] | None = None,
    delivery_fee_enabled: bool = False,
    invoice_footer_text: str | None = None,
    source_workbook_path: Path | None = None,
) -> None:
    _build_document_workbook(
        output_path,
        "Rechnung",
        customer_name,
        document_number,
        line_items,
        document_date,
        deposit_returns or [],
        delivery_fee_enabled,
        None,
        invoice_footer_text,
        source_workbook_path,
    )


def build_delivery_note_workbook(
    output_path: Path,
    customer_name: str,
    document_number: str,
    line_items: list[dict],
    document_date: str | None = None,
    deposit_returns: list[dict] | None = None,
    delivery_fee_enabled: bool = False,
    delivery_comment: str | None = None,
    footer_text: str | None = None,
    source_workbook_path: Path | None = None,
) -> None:
    _build_document_workbook(
        output_path,
        "Lieferschein",
        customer_name,
        document_number,
        line_items,
        document_date,
        deposit_returns or [],
        delivery_fee_enabled,
        delivery_comment,
        footer_text or DELIVERY_NOTE_FOOTER,
        source_workbook_path,
    )


def _build_document_workbook(
    output_path: Path,
    document_title: str,
    customer_name: str,
    document_number: str,
    line_items: list[dict],
    document_date: str | None,
    deposit_returns: list[dict],
    delivery_fee_enabled: bool,
    delivery_comment: str | None,
    footer_text: str | None,
    source_workbook_path: Path | None,
) -> None:
    ensure_parent_folder(output_path)

    preserve_customer_template = source_workbook_path is not None and source_workbook_path.exists()
    workbook = load_workbook(source_workbook_path if preserve_customer_template else TEMPLATE_PATH)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.forceFullCalc = True
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.calcOnSave = True
    sheet = workbook.active

    sheet["A8"] = document_title
    sheet["F8"] = document_number
    sheet["B10"] = customer_name
    sheet["C5"] = _document_date_value(document_date)
    sheet["D1"] = delivery_comment or None
    sheet["A45"] = footer_text or None
    sheet["B48"] = THANK_YOU_TEXT

    if len(line_items) > MAX_ITEM_ROW - FIRST_ITEM_ROW + 1:
        raise ValueError("Die Winklmeier-Vorlage erlaubt maximal 18 Positionszeilen.")
    if len(deposit_returns) > MAX_RETURN_ROW - FIRST_RETURN_ROW + 1:
        raise ValueError("Die Winklmeier-Vorlage erlaubt maximal 6 Pfandrueckgabe-Zeilen.")

    beleg_summen = berechne_beleg(
        positionen=[
            Position(
                bezeichnung=item["name"],
                menge=item["quantity"],
                preis_cents=item["unit_price_cents"],
                pfand_cents=item.get("deposit_cents", 0),
            )
            for item in line_items
        ],
        ruecknahmen=[
            PfandRueckgabe(
                bezeichnung=deposit_return["name"],
                menge=deposit_return["quantity"],
                pfand_cents=deposit_return["deposit_cents"],
            )
            for deposit_return in deposit_returns
        ],
        parameter=BelegParameter(lieferpauschale_aktiv=delivery_fee_enabled),
    )

    cached_formula_values = {}
    item_rows_by_name: dict[str, int] = {}
    empty_item_rows: list[int] = []
    for row in range(FIRST_ITEM_ROW, MAX_ITEM_ROW + 1):
        current_name = sheet.cell(row=row, column=2).value
        if preserve_customer_template and current_name:
            item_rows_by_name[_normalize_lookup(str(current_name))] = row
            sheet.cell(row=row, column=1).value = None
        else:
            empty_item_rows.append(row)
            sheet.cell(row=row, column=1).value = None
            sheet.cell(row=row, column=2).value = None
            sheet.cell(row=row, column=3).value = None
            sheet.cell(row=row, column=4).value = None
        sheet.cell(row=row, column=5, value=f"=(C{row}+D{row})*A{row}")
        cached_formula_values[f"E{row}"] = 0

    return_rows_by_deposit: dict[int, int] = {}
    empty_return_rows: list[int] = []
    for row in range(FIRST_RETURN_ROW, MAX_RETURN_ROW + 1):
        deposit_cents = _negative_money_cell_to_positive_cents(sheet.cell(row=row, column=3).value)
        if preserve_customer_template and deposit_cents:
            return_rows_by_deposit[deposit_cents] = row
            sheet.cell(row=row, column=1).value = None
        else:
            empty_return_rows.append(row)
            sheet.cell(row=row, column=1).value = None
            sheet.cell(row=row, column=2).value = None
            sheet.cell(row=row, column=3).value = None
        sheet.cell(row=row, column=5, value=f"=(C{row}+D{row})*A{row}")
        cached_formula_values[f"E{row}"] = 0

    for index, item in enumerate(line_items):
        lookup_name = _normalize_lookup(item["name"])
        target_row = item_rows_by_name.pop(lookup_name, None)
        if target_row is None:
            target_row = empty_item_rows.pop(0)
        quantity = item["quantity"]
        unit_price = item["unit_price_cents"] / 100
        deposit = item.get("deposit_cents", 0) / 100
        line_total_cents = beleg_summen.positionssummen_cents[index]
        sheet.cell(row=target_row, column=1, value=quantity if quantity > 0 else None)
        sheet.cell(row=target_row, column=2, value=item["name"])
        sheet.cell(row=target_row, column=3, value=deposit)
        sheet.cell(row=target_row, column=4, value=unit_price)
        sheet.cell(row=target_row, column=5, value=f"=(C{target_row}+D{target_row})*A{target_row}")
        cached_formula_values[f"E{target_row}"] = _cents_to_euro(line_total_cents)

    sheet["A31"] = 1 if delivery_fee_enabled else 0
    sheet["B31"] = "Lieferpauschale"
    sheet["C31"] = None
    sheet["D31"] = _cents_to_euro(DELIVERY_FEE_CENTS)
    sheet["E31"] = "=(C31+D31)*A31"
    cached_formula_values["E31"] = _cents_to_euro(beleg_summen.lieferpauschale_cents)

    for index, deposit_return in enumerate(deposit_returns):
        quantity = deposit_return["quantity"]
        deposit_cents = deposit_return["deposit_cents"]
        target_row = return_rows_by_deposit.pop(deposit_cents, None)
        if target_row is None:
            target_row = empty_return_rows.pop(0)
        line_total_cents = beleg_summen.ruecknahme_summen_cents[index]
        sheet.cell(row=target_row, column=1, value=quantity)
        if not preserve_customer_template:
            sheet.cell(row=target_row, column=2, value=deposit_return["name"])
        sheet.cell(row=target_row, column=3, value=-deposit_cents / 100)
        sheet.cell(row=target_row, column=5, value=f"=(C{target_row}+D{target_row})*A{target_row}")
        cached_formula_values[f"E{target_row}"] = _cents_to_euro(line_total_cents)

    sheet["A32"] = f"=SUM(A{FIRST_ITEM_ROW}:A{MAX_ITEM_ROW})"
    sheet["F32"] = f"=SUM(E{FIRST_ITEM_ROW}:E31)"
    sheet["F40"] = f"=SUM(E{FIRST_RETURN_ROW}:E{MAX_RETURN_ROW})"
    sheet["F41"] = "=ROUND(F43/1.19,2)"
    sheet["F42"] = "=F43-F41"
    sheet["F43"] = "=F32+F40"

    cached_formula_values.update(
        {
            "A32": beleg_summen.mengen_summe,
            "F32": _cents_to_euro(beleg_summen.lieferwert_cents + beleg_summen.lieferpauschale_cents),
            "F40": _cents_to_euro(beleg_summen.pfand_rueckgabe_cents),
            "F41": _cents_to_euro(beleg_summen.netto_cents),
            "F42": _cents_to_euro(beleg_summen.mwst_cents),
            "F43": _cents_to_euro(beleg_summen.brutto_cents),
        }
    )

    workbook.save(output_path)
    _store_cached_formula_values(output_path, cached_formula_values)


def _document_date_value(document_date: str | None) -> datetime:
    if document_date:
        try:
            return datetime.strptime(document_date, "%Y-%m-%d")
        except ValueError:
            pass
    return datetime.combine(date.today(), datetime.min.time())


def _cents_to_euro(cents: int) -> int | float:
    euros = cents / 100
    return int(euros) if cents % 100 == 0 else euros


def _normalize_lookup(value: str) -> str:
    return " ".join(value.strip().casefold().split())


def _negative_money_cell_to_positive_cents(value) -> int:
    if value is None:
        return 0
    try:
        cents = round(float(value) * 100)
    except (TypeError, ValueError):
        return 0
    return abs(cents) if cents < 0 else 0


def _store_cached_formula_values(output_path: Path, cached_formula_values: dict[str, int | float]) -> None:
    worksheet_path = "xl/worksheets/sheet1.xml"
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)

    with NamedTemporaryFile(delete=False, dir=output_path.parent, suffix=".xlsx") as tmp_file:
        tmp_path = Path(tmp_file.name)
    try:
        with ZipFile(output_path, "r") as source_zip:
            with ZipFile(tmp_path, "w", ZIP_DEFLATED) as target_zip:
                for item in source_zip.infolist():
                    data = source_zip.read(item.filename)
                    if item.filename == worksheet_path:
                        data = _worksheet_xml_with_cached_values(data, cached_formula_values, namespace)
                    target_zip.writestr(item, data)
        tmp_path.replace(output_path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def _worksheet_xml_with_cached_values(
    worksheet_xml: bytes,
    cached_formula_values: dict[str, int | float],
    namespace: str,
) -> bytes:
    root = ET.fromstring(worksheet_xml)
    ns = f"{{{namespace}}}"
    for cell in root.findall(f".//{ns}c"):
        coordinate = cell.attrib.get("r")
        if coordinate not in cached_formula_values:
            continue
        if cell.find(f"{ns}f") is None:
            continue
        cell.attrib.pop("t", None)
        value = cell.find(f"{ns}v")
        if value is None:
            value = ET.SubElement(cell, f"{ns}v")
        value.text = str(cached_formula_values[coordinate])
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
