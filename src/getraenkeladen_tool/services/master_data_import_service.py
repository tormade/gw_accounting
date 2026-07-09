from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Customer, MasterDataChange, Product
from ..schemas import CustomerCreate, ProductCreate
from .customer_service import create_customer, update_customer
from .product_service import create_product, update_product


ARTICLE_FILE_NAME = "Artikel Liste Preise.xlsx"
CUSTOMER_FILE_NAME = "Lieferkunden Liste.xlsx"


@dataclass(slots=True)
class MasterDataImportResult:
    products_created: int = 0
    products_updated: int = 0
    customers_created: int = 0
    customers_updated: int = 0
    products: list[Product] = field(default_factory=list)
    customers: list[Customer] = field(default_factory=list)
    changes: list[MasterDataChange] = field(default_factory=list)


@dataclass(frozen=True, slots=True)
class MasterDataPreviewItem:
    entity_type: str
    action: str
    name: str
    source_row: int
    changes: tuple[tuple[str, str | None, str | None], ...] = ()
    message: str | None = None


@dataclass(frozen=True, slots=True)
class MasterDataImportPreview:
    items: tuple[MasterDataPreviewItem, ...]
    missing_required_files: tuple[str, ...] = ()

    @property
    def can_import(self) -> bool:
        return not self.missing_required_files

    @property
    def products_created(self) -> int:
        return self._count("product", "create")

    @property
    def products_updated(self) -> int:
        return self._count("product", "update")

    @property
    def customers_created(self) -> int:
        return self._count("customer", "create")

    @property
    def customers_updated(self) -> int:
        return self._count("customer", "update")

    @property
    def unclear_rows(self) -> int:
        return self._count_any("skip")

    @property
    def summary_text(self) -> str:
        return (
            f"{self.products_created} Produkt neu, {self.products_updated} Produkte aktualisiert, "
            f"{self.customers_created} Kunden neu, {self.customers_updated} Kunden aktualisiert, "
            f"{self.unclear_rows} Zeilen unklar."
        )

    @property
    def safety_groups(self) -> dict[str, int]:
        return {
            "neu": self._count_any("create"),
            "geaendert": self._count_any("update"),
            "unsicher": self._count_any("unclear"),
            "uebersprungen": self._count_any("skip"),
        }

    @property
    def safety_report_text(self) -> str:
        groups = self.safety_groups
        missing_files = "\n".join(f"Fehlt: {file_name}" for file_name in self.missing_required_files)
        required_files = f"\n\nBenötigte Dateien fehlen:\n{missing_files}" if missing_files else ""
        return (
            "Import-Sicherheitspruefung:\n"
            f"Neu: {groups['neu']}\n"
            f"Geaendert: {groups['geaendert']}\n"
            f"Unsicher: {groups['unsicher']}\n"
            f"Uebersprungen: {groups['uebersprungen']}\n\n"
            f"{self.summary_text}{required_files}"
        )

    def _count(self, entity_type: str, action: str) -> int:
        return sum(1 for item in self.items if item.entity_type == entity_type and item.action == action)

    def _count_any(self, action: str) -> int:
        return sum(1 for item in self.items if item.action == action)


def preview_master_data_from_folder(session: Session, input_dir: Path) -> MasterDataImportPreview:
    items: list[MasterDataPreviewItem] = []
    article_path = input_dir / ARTICLE_FILE_NAME
    customer_path = input_dir / CUSTOMER_FILE_NAME
    if article_path.exists():
        items.extend(_preview_products(session, article_path))
    if customer_path.exists():
        items.extend(_preview_customers(session, customer_path))
    missing_required_files = tuple(
        file_name
        for file_name, path in ((ARTICLE_FILE_NAME, article_path), (CUSTOMER_FILE_NAME, customer_path))
        if not path.exists()
    )
    return MasterDataImportPreview(tuple(items), missing_required_files)


def import_master_data_from_folder(session: Session, input_dir: Path) -> MasterDataImportResult:
    preview = preview_master_data_from_folder(session, input_dir)
    if not preview.can_import:
        missing_files = ", ".join(preview.missing_required_files)
        raise ValueError(f"Import nicht möglich. Benötigte Dateien fehlen: {missing_files}.")
    result = MasterDataImportResult()
    first_new_change_id = (session.scalar(select(func.max(MasterDataChange.id))) or 0) + 1
    article_path = input_dir / ARTICLE_FILE_NAME
    customer_path = input_dir / CUSTOMER_FILE_NAME
    if article_path.exists():
        _import_products(session, article_path, result)
    if customer_path.exists():
        _import_customers(session, customer_path, result)
    session.commit()
    result.changes = list(
        session.scalars(
            select(MasterDataChange)
            .where(MasterDataChange.id >= first_new_change_id)
            .order_by(MasterDataChange.id)
        )
    )
    return result


def _preview_products(session: Session, path: Path) -> list[MasterDataPreviewItem]:
    items: list[MasterDataPreviewItem] = []
    for row_number, row in _sheet_rows(path):
        name = _clean_text(row.get("Prudukt Bezeichnung"))
        if not name:
            items.append(MasterDataPreviewItem("product", "skip", "", row_number, message="Produktname fehlt."))
            continue
        payload = _product_payload_from_row(path, row_number, row, name)
        existing = _find_product_by_name(session, name)
        if existing is None:
            items.append(MasterDataPreviewItem("product", "create", name, row_number))
        else:
            changes = _product_preview_changes(existing, payload)
            if changes:
                items.append(MasterDataPreviewItem("product", "update", name, row_number, changes))
    return items


def _preview_customers(session: Session, path: Path) -> list[MasterDataPreviewItem]:
    items: list[MasterDataPreviewItem] = []
    for row_number, row in _sheet_rows(path):
        name = _clean_text(row.get("Name"))
        if not name:
            items.append(MasterDataPreviewItem("customer", "skip", "", row_number, message="Kundenname fehlt."))
            continue
        payload = _customer_payload_from_row(path, row_number, row, name)
        existing = _find_customer_by_name(session, name)
        if existing is None:
            items.append(MasterDataPreviewItem("customer", "create", name, row_number))
        else:
            changes = _customer_preview_changes(existing, payload)
            if changes:
                items.append(MasterDataPreviewItem("customer", "update", name, row_number, changes))
    return items


def _import_products(session: Session, path: Path, result: MasterDataImportResult) -> None:
    for row_number, row in _sheet_rows(path):
        name = _clean_text(row.get("Prudukt Bezeichnung"))
        if not name:
            continue
        payload = _product_payload_from_row(path, row_number, row, name)
        existing = _find_product_by_name(session, name)
        if existing is None:
            product = create_product(session, payload)
            result.products_created += 1
        else:
            payload.is_active = existing.is_active
            product = update_product(session, existing.id, payload)
            result.products_updated += 1
        result.products.append(product)


def _import_customers(session: Session, path: Path, result: MasterDataImportResult) -> None:
    for row_number, row in _sheet_rows(path):
        name = _clean_text(row.get("Name"))
        if not name:
            continue
        payload = _customer_payload_from_row(path, row_number, row, name)
        existing = _find_customer_by_name(session, name)
        if existing is None:
            customer = create_customer(session, payload)
            result.customers_created += 1
        else:
            payload.is_active = existing.is_active
            customer = update_customer(session, existing.id, payload)
            result.customers_updated += 1
        result.customers.append(customer)


def _product_payload_from_row(path: Path, row_number: int, row: dict, name: str) -> ProductCreate:
    return ProductCreate(
        name=name,
        unit=_unit_from_article_row(row),
        standard_price_cents=_price_to_cents(row.get("Liefer Preis") or row.get("Laden VK") or row.get("Re.-Preis")),
        default_deposit_cents=_price_to_cents(row.get("Pfand")),
        source_file=str(path),
        source_row=row_number,
    )


def _customer_payload_from_row(path: Path, row_number: int, row: dict, name: str) -> CustomerCreate:
    return CustomerCreate(
        name=name,
        folder_path=f"Kunden/{name}",
        address=_address_from_customer_row(row),
        contact_name=_clean_text(row.get("Ansprechpartner")),
        contact_email=_clean_text(row.get("e-mail (Kontakt)")) or _clean_text(row.get("e-mail (Re Versand)")),
        next_contact_date=_date_to_iso(row.get("nächster Kontakt")),
        delivery_notes=_delivery_notes_from_customer_row(row),
        opening_hours=_opening_hours_from_customer_row(row),
        internal_notes=_clean_text(row.get("ABO ")),
        source_file=str(path),
        source_row=row_number,
    )


def _product_preview_changes(product: Product, payload: ProductCreate) -> tuple[tuple[str, str | None, str | None], ...]:
    return _preview_changes(
        product,
        payload,
        ("unit", "standard_price_cents", "default_deposit_cents", "source_file", "source_row"),
    )


def _customer_preview_changes(customer: Customer, payload: CustomerCreate) -> tuple[tuple[str, str | None, str | None], ...]:
    return _preview_changes(
        customer,
        payload,
        (
            "folder_path",
            "address",
            "contact_name",
            "contact_email",
            "next_contact_date",
            "delivery_notes",
            "opening_hours",
            "internal_notes",
            "source_file",
            "source_row",
        ),
    )


def _preview_changes(entity, payload, fields: tuple[str, ...]) -> tuple[tuple[str, str | None, str | None], ...]:
    changes = []
    for field_name in fields:
        old_value = getattr(entity, field_name)
        new_value = getattr(payload, field_name)
        if old_value != new_value:
            changes.append((field_name, _string_or_none(old_value), _string_or_none(new_value)))
    return tuple(changes)


def _string_or_none(value) -> str | None:
    return None if value is None else str(value)


def _sheet_rows(path: Path):
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook["Tabelle1"] if "Tabelle1" in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return
    headers = [_clean_text(value) for value in rows[0]]
    for row_number, values in enumerate(rows[1:], start=2):
        yield row_number, {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header}


def _find_product_by_name(session: Session, name: str) -> Product | None:
    return session.scalar(select(Product).where(Product.name == name))


def _find_customer_by_name(session: Session, name: str) -> Customer | None:
    return session.scalar(select(Customer).where(Customer.name == name))


def _unit_from_article_row(row: dict) -> str:
    bottle_count = _clean_text(row.get("Fl."))
    amount = _clean_text(row.get("Menge"))
    if bottle_count and amount:
        return f"Kiste {bottle_count}x{amount}"
    return "Kiste"


def _address_from_customer_row(row: dict) -> str | None:
    street = _clean_text(row.get("Straße"))
    house_number = _clean_text(row.get("Hausnummer"))
    postal_code = _clean_text(row.get("PLZ"))
    city = _clean_text(row.get("Ort"))
    street_part = " ".join(part for part in (street, house_number) if part)
    city_part = " ".join(part for part in (postal_code, city) if part)
    return ", ".join(part for part in (street_part, city_part) if part) or None


def _delivery_notes_from_customer_row(row: dict) -> str | None:
    notes = [
        _clean_text(row.get("Bemerkungen")),
        _clean_text(row.get("ABO ")),
        _clean_text(row.get(None)),
    ]
    if _clean_text(row.get("Vormittag")):
        notes.append("vormittag")
    return " | ".join(note for note in notes if note) or None


def _opening_hours_from_customer_row(row: dict) -> str | None:
    if _clean_text(row.get("Vormittag")):
        return "vormittag"
    return None


def _date_to_iso(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _clean_text(value)


def _price_to_cents(value) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, str):
        normalized = value.replace(".", "").replace(",", ".").strip()
        if not normalized:
            return 0
        return int(round(float(normalized) * 100))
    return int(round(float(value) * 100))


def _clean_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
