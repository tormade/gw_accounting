from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload
from openpyxl import load_workbook

from ..models import Document
from .pdf_service import build_document_pdf


def list_documents_by_type(session: Session, document_type: str) -> list[Document]:
    return list(
        session.scalars(
            select(Document)
            .options(selectinload(Document.customer))
            .where(Document.document_type == document_type)
            .where(Document.number_released == False)  # noqa: E712
            .order_by(Document.delivery_date.desc(), Document.id.desc())
        )
    )


def regenerate_document_asset(session: Session, document_id: int, asset: str) -> Document:
    if asset not in {"excel", "pdf"}:
        raise ValueError("Es kann nur Excel oder PDF neu erzeugt werden.")

    document = session.get(Document, document_id)
    if document is None or document.number_released:
        raise ValueError("Der Beleg wurde nicht gefunden.")
    if asset == "excel":
        raise ValueError(
            "Die Excel-Datei kann nicht sicher automatisch neu erzeugt werden. "
            "Bitte den Beleg im Kundenordner pruefen oder aus dem Auftrag neu erstellen."
        )

    excel_path = Path(document.excel_path)
    if not excel_path.exists():
        raise ValueError("Die PDF kann nur neu erzeugt werden, wenn die Excel-Datei noch vorhanden ist.")

    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook.active
    line_items = _line_items_from_workbook(sheet)
    deposit_returns = _deposit_returns_from_workbook(sheet)
    build_document_pdf(
        Path(document.pdf_path),
        document.document_type,
        document.customer.name,
        document.document_number,
        line_items,
        deposit_returns=deposit_returns,
        document_date=document.delivery_date,
        customer_address=document.customer.address,
        delivery_fee_enabled=bool(sheet["A31"].value),
        note_text=document.delivery_comment,
        footer_text=document.footer_text,
    )
    return document


def _line_items_from_workbook(sheet) -> list[dict]:
    items = []
    for row in range(13, 31):
        name = sheet.cell(row=row, column=2).value
        quantity = sheet.cell(row=row, column=1).value
        if not name or not quantity:
            continue
        items.append(
            {
                "name": str(name),
                "quantity": int(quantity),
                "deposit_cents": _euro_to_cents(sheet.cell(row=row, column=3).value),
                "unit_price_cents": _euro_to_cents(sheet.cell(row=row, column=4).value),
            }
        )
    return items


def _deposit_returns_from_workbook(sheet) -> list[dict]:
    returns = []
    for row in range(35, 43):
        name = sheet.cell(row=row, column=2).value
        quantity = sheet.cell(row=row, column=1).value
        deposit_value = sheet.cell(row=row, column=3).value
        if not name or not quantity or not deposit_value:
            continue
        returns.append(
            {
                "name": str(name),
                "quantity": int(quantity),
                "deposit_cents": abs(_euro_to_cents(deposit_value)),
            }
        )
    return returns


def _euro_to_cents(value) -> int:
    if value is None:
        return 0
    return int(round(float(value) * 100))
