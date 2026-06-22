from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.orm import Session

from ..models import Document, NumberSequence, OpenItem, Order


@dataclass(frozen=True)
class NumberSuggestions:
    order_number: str
    delivery_note_number: str
    invoice_number: str


@dataclass(frozen=True)
class NumberSequenceStatus:
    sequence_key: str
    label: str
    prefix: str
    next_number: str


@dataclass(frozen=True)
class NumberConflict:
    sequence_key: str
    number: str
    kind: str
    customer_name: str
    date: str
    detail: str


SEQUENCE_DEFINITIONS = {
    "order": {"label": "Auftrag", "prefix": "AUF", "start": 1001},
    "delivery_note": {"label": "Lieferschein", "prefix": "LS", "start": 3001},
    "invoice": {"label": "Rechnung", "prefix": "RG", "start": 3001},
}
NUMBER_SEQUENCE_WORKBOOK = "Nummernkreise.xlsx"
NUMBER_SEQUENCE_SHEET = "Nummernkreise"
NUMBER_SEQUENCE_HEADERS = ("Bereich", "Schluessel", "Praefix", "Naechste Nummer", "Hinweis")


def suggest_next_numbers(session: Session) -> NumberSuggestions:
    return NumberSuggestions(
        order_number=next_order_number(session),
        delivery_note_number=next_document_number(
            session,
            document_type="Lieferschein",
            prefix="LS",
            start=3001,
            aliases=("Lieferauftrag",),
        ),
        invoice_number=next_document_number(session, document_type="Rechnung", prefix="RG", start=3001),
    )


def next_order_number(session: Session) -> str:
    existing_numbers = session.scalars(
        select(Order.order_number)
        .where(Order.status != "archiviert")
        .where(Order.number_released == False)  # noqa: E712
    ).all()
    return _next_number_with_manual_floor(
        session,
        sequence_key="order",
        existing_numbers=existing_numbers,
        prefix="AUF",
        start=1001,
    )


def next_document_number(
    session: Session,
    document_type: str,
    prefix: str,
    start: int,
    aliases: tuple[str, ...] = (),
) -> str:
    document_types = (document_type, *aliases)
    existing_numbers = session.scalars(
        select(Document.document_number).where(Document.document_type.in_(document_types))
        .where(Document.number_released == False)  # noqa: E712
    ).all()
    sequence_key = _sequence_key_for_prefix(prefix)
    return _next_number_with_manual_floor(
        session,
        sequence_key=sequence_key,
        existing_numbers=existing_numbers,
        prefix=prefix,
        start=start,
    )


def list_number_sequence_statuses(session: Session) -> list[NumberSequenceStatus]:
    suggestions = suggest_next_numbers(session)
    next_numbers = {
        "order": suggestions.order_number,
        "delivery_note": suggestions.delivery_note_number,
        "invoice": suggestions.invoice_number,
    }
    statuses = []
    for sequence_key, definition in SEQUENCE_DEFINITIONS.items():
        prefix = str(definition["prefix"])
        sequence = session.scalar(
            select(NumberSequence).where(NumberSequence.sequence_key == sequence_key).limit(1)
        )
        next_number = f"{prefix}-{sequence.next_number}" if sequence is not None else next_numbers[sequence_key]
        statuses.append(
            NumberSequenceStatus(
                sequence_key=sequence_key,
                label=str(definition["label"]),
                prefix=prefix,
                next_number=next_number,
            )
        )
    return statuses


def set_next_number(session: Session, sequence_key: str, prefix: str, value: str) -> NumberSequence:
    clean_value = _normalize_sequence_value(value, prefix)
    match = re.match(rf"^{re.escape(prefix)}-(\d+)$", clean_value)
    next_number = int(match.group(1))
    sequence = session.scalar(
        select(NumberSequence).where(NumberSequence.sequence_key == sequence_key).limit(1)
    )
    if sequence is None:
        sequence = NumberSequence(sequence_key=sequence_key, prefix=prefix, next_number=next_number)
        session.add(sequence)
    else:
        sequence.prefix = prefix
        sequence.next_number = next_number
    session.commit()
    session.refresh(sequence)
    return sequence


def reset_number_sequences_to_defaults(session: Session) -> list[NumberSequence]:
    sequences = []
    for sequence_key, definition in SEQUENCE_DEFINITIONS.items():
        sequences.append(
            set_next_number(
                session,
                sequence_key=sequence_key,
                prefix=str(definition["prefix"]),
                value=f"{definition['prefix']}-{definition['start']}",
            )
        )
    return sequences


def number_sequence_workbook_path(session: Session) -> Path:
    database = getattr(session.bind.url, "database", None)
    if database in (None, "", ":memory:"):
        return Path.cwd() / NUMBER_SEQUENCE_WORKBOOK
    database_path = Path(database)
    return database_path.parent.parent / NUMBER_SEQUENCE_WORKBOOK


def ensure_number_sequence_workbook(session: Session, workbook_path: Path | None = None) -> Path:
    path = workbook_path or number_sequence_workbook_path(session)
    if not path.exists():
        write_number_sequences_to_workbook(session, path)
    return path


def load_number_sequences_from_workbook(session: Session, workbook_path: Path | None = None) -> Path:
    path = ensure_number_sequence_workbook(session, workbook_path)
    workbook = load_workbook(path)
    if NUMBER_SEQUENCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Die Datei {path.name} braucht ein Blatt '{NUMBER_SEQUENCE_SHEET}'.")
    sheet = workbook[NUMBER_SEQUENCE_SHEET]
    headers = [sheet.cell(row=1, column=column).value for column in range(1, len(NUMBER_SEQUENCE_HEADERS) + 1)]
    if tuple(headers) != NUMBER_SEQUENCE_HEADERS:
        raise ValueError("Nummernkreise.xlsx hat nicht die erwarteten Spalten.")
    for row in range(2, sheet.max_row + 1):
        sequence_key = str(sheet.cell(row=row, column=2).value or "").strip()
        prefix = str(sheet.cell(row=row, column=3).value or "").strip()
        next_number = str(sheet.cell(row=row, column=4).value or "").strip()
        if sequence_key in SEQUENCE_DEFINITIONS and prefix and next_number:
            set_next_number(session, sequence_key=sequence_key, prefix=prefix, value=next_number)
    return path


def write_number_sequences_to_workbook(session: Session, workbook_path: Path | None = None) -> Path:
    path = workbook_path or number_sequence_workbook_path(session)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = NUMBER_SEQUENCE_SHEET
    sheet.append(NUMBER_SEQUENCE_HEADERS)
    statuses = list_number_sequence_statuses(session)
    for status in statuses:
        sheet.append(
            [
                status.label,
                status.sequence_key,
                status.prefix,
                status.next_number,
                "Diese Nummer kann in Excel geaendert werden. Danach Einstellungen in der App neu laden.",
            ]
        )
    for column_width, column in ((18, "A"), (18, "B"), (10, "C"), (18, "D"), (78, "E")):
        sheet.column_dimensions[column].width = column_width
    workbook.save(path)
    return path


def check_number_conflict(session: Session, sequence_key: str, value: str) -> NumberConflict | None:
    definition = SEQUENCE_DEFINITIONS[sequence_key]
    prefix = str(definition["prefix"])
    clean_value = _normalize_sequence_value(value, prefix)
    if sequence_key == "order":
        order = session.scalar(
            select(Order)
            .options(selectinload(Order.customer))
            .where(Order.order_number == clean_value)
            .where(Order.status != "archiviert")
            .where(Order.number_released == False)  # noqa: E712
            .limit(1)
        )
        if order is None:
            return None
        return NumberConflict(
            sequence_key=sequence_key,
            number=clean_value,
            kind="Auftrag",
            customer_name=order.customer.name if order.customer is not None else "",
            date=order.delivery_date,
            detail=f"Auftrag {order.order_number}",
        )

    document_types = _document_types_for_sequence(sequence_key)
    document = session.scalar(
        select(Document)
        .options(selectinload(Document.customer))
        .where(Document.document_number == clean_value)
        .where(Document.document_type.in_(document_types))
        .where(Document.number_released == False)  # noqa: E712
        .limit(1)
    )
    if document is None:
        return None
    return NumberConflict(
        sequence_key=sequence_key,
        number=clean_value,
        kind=document.document_type,
        customer_name=document.customer.name if document.customer is not None else "",
        date=document.delivery_date or "",
        detail=f"{document.document_type} {document.document_number}",
    )


def release_number(session: Session, sequence_key: str, value: str) -> NumberConflict | None:
    conflict = check_number_conflict(session, sequence_key, value)
    if conflict is None:
        return None
    if sequence_key == "order":
        order = session.scalar(
            select(Order)
            .where(Order.order_number == conflict.number)
            .where(Order.status != "archiviert")
            .where(Order.number_released == False)  # noqa: E712
            .limit(1)
        )
        if order is not None:
            order.number_released = True
    else:
        document = session.scalar(
            select(Document)
            .where(Document.document_number == conflict.number)
            .where(Document.document_type.in_(_document_types_for_sequence(sequence_key)))
            .where(Document.number_released == False)  # noqa: E712
            .limit(1)
        )
        if document is not None:
            document.number_released = True
            open_item = session.scalar(select(OpenItem).where(OpenItem.document_id == document.id).limit(1))
            if open_item is not None and open_item.status == "offen":
                open_item.status = "freigegeben"
    session.commit()
    return conflict


def next_number_for_prefix(existing_numbers: list[str], prefix: str, start: int) -> str:
    highest = start - 1
    pattern = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    for number in existing_numbers:
        match = pattern.match(number.strip())
        if match is None:
            continue
        highest = max(highest, int(match.group(1)))
    return f"{prefix}-{highest + 1}"


def _next_number_with_manual_floor(
    session: Session,
    sequence_key: str,
    existing_numbers: list[str],
    prefix: str,
    start: int,
) -> str:
    sequence = session.scalar(
        select(NumberSequence).where(NumberSequence.sequence_key == sequence_key).limit(1)
    )
    if sequence is None:
        return next_number_for_prefix(existing_numbers, prefix=prefix, start=start)
    return next_available_number_from(existing_numbers, prefix=prefix, start=sequence.next_number)


def next_available_number_from(existing_numbers: list[str], prefix: str, start: int) -> str:
    used_numbers = _used_number_values(existing_numbers, prefix)
    candidate = start
    while candidate in used_numbers:
        candidate += 1
    return f"{prefix}-{candidate}"


def _used_number_values(existing_numbers: list[str], prefix: str) -> set[int]:
    pattern = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    used_numbers = set()
    for number in existing_numbers:
        match = pattern.match(number.strip())
        if match is not None:
            used_numbers.add(int(match.group(1)))
    return used_numbers


def _sequence_key_for_prefix(prefix: str) -> str:
    for sequence_key, definition in SEQUENCE_DEFINITIONS.items():
        if definition["prefix"] == prefix:
            return sequence_key
    return prefix.lower()


def _normalize_sequence_value(value: str, prefix: str) -> str:
    clean_value = value.strip().upper()
    pattern = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    match = pattern.match(clean_value)
    if match is None:
        raise ValueError(f"Nummer muss dem Format {prefix}-1234 entsprechen.")
    return clean_value


def _document_types_for_sequence(sequence_key: str) -> tuple[str, ...]:
    if sequence_key == "delivery_note":
        return ("Lieferschein", "Lieferauftrag")
    if sequence_key == "invoice":
        return ("Rechnung",)
    return ()
