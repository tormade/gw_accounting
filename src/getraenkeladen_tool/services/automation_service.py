from dataclasses import dataclass
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..models import Customer, CustomerAssortmentItem, Document, DocumentLine, OpenItem, OnboardingIssue, Order, Product
from ..schemas import OrderCreate


@dataclass(frozen=True, slots=True)
class CustomerQuickstart:
    customer_id: int
    customer_name: str
    folder_exists: bool
    latest_order_number: str | None
    open_invoice_numbers: tuple[str, ...]
    suggestions: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class DuplicateRisk:
    entity_type: str
    existing_label: str
    message: str


@dataclass(frozen=True, slots=True)
class AutomationWarning:
    code: str
    message: str
    severity: str = "warning"


@dataclass(frozen=True, slots=True)
class AssetCheck:
    code: str
    message: str
    ok: bool


@dataclass(frozen=True, slots=True)
class DocumentAssetVerification:
    document_number: str
    ok: bool
    checks: tuple[AssetCheck, ...]


@dataclass(frozen=True, slots=True)
class PaymentHint:
    due_date: str | None
    footer_text: str | None


@dataclass(frozen=True, slots=True)
class MonthEndCheck:
    code: str
    message: str
    severity: str = "warning"


@dataclass(frozen=True, slots=True)
class TypicalOrderLine:
    product_id: int | None
    product_name: str
    suggested_quantity: int
    seen_in_orders: int
    last_quantity: int
    min_quantity: int
    max_quantity: int
    frequency_label: str


@dataclass(frozen=True, slots=True)
class WorkCockpitItem:
    code: str
    title: str
    detail: str
    severity: str = "warning"


@dataclass(frozen=True, slots=True)
class CustomerFolderExcelPreview:
    path: Path
    document_number: str | None
    document_date: str | None
    filled_quantity_rows: int


@dataclass(frozen=True, slots=True)
class DocumentSequenceIssue:
    code: str
    order_number: str
    message: str
    severity: str = "warning"


@dataclass(frozen=True, slots=True)
class CustomerPriceDeviation:
    customer_id: int
    product_id: int
    product_name: str
    customer_price_cents: int
    master_price_cents: int
    decision: str


def get_customer_quickstart(session: Session, customer_id: int) -> CustomerQuickstart:
    customer = _customer_or_error(session, customer_id)
    latest_order = session.scalar(
        select(Order)
        .where(Order.customer_id == customer.id)
        .where(Order.status != "archiviert")
        .order_by(Order.delivery_date.desc(), Order.id.desc())
        .limit(1)
    )
    open_items = list(
        session.scalars(
            select(OpenItem)
            .where(OpenItem.customer_name == customer.name)
            .where(OpenItem.status == "offen")
            .order_by(OpenItem.due_date, OpenItem.document_number)
        )
    )
    folder_exists = Path(customer.folder_path).is_dir()
    suggestions = ["Neue Bestellung aus letzter Kunden-Excel starten"]
    if latest_order is not None:
        suggestions.append(f"Letzte Bestellung {latest_order.order_number} oeffnen")
    if open_items:
        suggestions.append(f"{len(open_items)} offene Rechnung(en) pruefen")
    if customer.delivery_notes:
        suggestions.append("Lieferhinweise vorlesen")
    if not folder_exists:
        suggestions.append("Kundenordner pruefen")
    return CustomerQuickstart(
        customer_id=customer.id,
        customer_name=customer.name,
        folder_exists=folder_exists,
        latest_order_number=latest_order.order_number if latest_order is not None else None,
        open_invoice_numbers=tuple(item.document_number for item in open_items),
        suggestions=tuple(suggestions),
    )


def find_duplicate_risks(
    session: Session,
    customer_name: str = "",
    product_name: str = "",
    document_number: str = "",
) -> list[DuplicateRisk]:
    risks: list[DuplicateRisk] = []
    if customer_name.strip():
        for customer in session.scalars(select(Customer).where(Customer.is_active == True)):  # noqa: E712
            if _looks_similar(customer_name, customer.name):
                risks.append(DuplicateRisk("customer", customer.name, f"Aehnlicher Kunde vorhanden: {customer.name}"))
    if product_name.strip():
        for product in session.scalars(select(Product).where(Product.is_active == True)):  # noqa: E712
            if _looks_similar(product_name, product.name):
                risks.append(DuplicateRisk("product", product.name, f"Aehnlicher Artikel vorhanden: {product.name}"))
    if document_number.strip():
        normalized_number = _compact(document_number)
        for document in session.scalars(select(Document).where(Document.number_released == False)):  # noqa: E712
            if _compact(document.document_number) == normalized_number:
                risks.append(DuplicateRisk("document", document.document_number, f"Belegnummer existiert bereits: {document.document_number}"))
    return risks


def validate_order_before_save(session: Session, payload: OrderCreate, exclude_order_id: int | None = None) -> list[AutomationWarning]:
    warnings: list[AutomationWarning] = []
    customer = session.get(Customer, payload.customer_id)
    if customer is None:
        warnings.append(AutomationWarning("missing_customer", "Kunde wurde nicht gefunden.", "error"))
    else:
        if not (customer.address or "").strip():
            warnings.append(AutomationWarning("missing_customer_address", "Beim Kunden fehlt die Adresse."))
        if not (customer.payment_method or "").strip():
            warnings.append(AutomationWarning("missing_customer_payment_method", "Beim Kunden fehlt die Zahlart."))
    if not any(line.quantity > 0 for line in payload.lines):
        warnings.append(AutomationWarning("no_positive_quantity", "Die Bestellung enthaelt keine Position mit Menge groesser 0.", "error"))
    order_number = payload.order_number.strip()
    duplicate_query = (
        select(Order.id)
        .where(Order.order_number == order_number)
        .where(Order.status != "archiviert")
        .where(Order.number_released == False)  # noqa: E712
        .limit(1)
    )
    if exclude_order_id is not None:
        duplicate_query = duplicate_query.where(Order.id != exclude_order_id)
    if session.scalar(duplicate_query) is not None:
        warnings.append(AutomationWarning("duplicate_order_number", "Auftragsnummer ist bereits vorhanden.", "error"))
    seen_products: set[int] = set()
    for line in payload.lines:
        product = session.get(Product, line.product_id)
        if product is None:
            warnings.append(AutomationWarning("missing_product", f"Produkt {line.product_id} wurde nicht gefunden.", "error"))
            continue
        if line.product_id in seen_products:
            warnings.append(AutomationWarning("duplicate_product", f"Artikel {product.name} ist doppelt in der Bestellung."))
        seen_products.add(line.product_id)
        if line.quantity >= 100:
            warnings.append(AutomationWarning("high_quantity", f"Ungewoehnlich hohe Menge bei {product.name}: {line.quantity}."))
        effective_price = product.standard_price_cents if line.unit_price_cents is None else line.unit_price_cents
        if effective_price == 0:
            warnings.append(AutomationWarning("zero_price", f"Preis bei {product.name} ist 0,00 EUR."))
        if product.default_deposit_cents and line.deposit_cents != product.default_deposit_cents:
            warnings.append(
                AutomationWarning(
                    "deposit_mismatch",
                    f"Pfand bei {product.name} weicht vom Artikelstamm ab.",
                )
            )
    return warnings


def verify_document_assets(
    session: Session,
    document_id: int,
    expected_assets: set[str] | None = None,
) -> DocumentAssetVerification:
    document = session.get(Document, document_id)
    if document is None:
        raise ValueError("Beleg wurde nicht gefunden.")
    assets = expected_assets or {"excel", "pdf"}
    checks = []
    if "excel" in assets:
        checks.append(_path_check("missing_excel", "Excel-Datei vorhanden.", document.excel_path))
    if "pdf" in assets:
        checks.append(_path_check("missing_pdf", "PDF-Datei vorhanden.", document.pdf_path))
    snapshot_total = session.scalar(
        select(DocumentLine.total_cents)
        .where(DocumentLine.document_id == document.id)
        .where(DocumentLine.line_type == "summary")
        .limit(1)
    )
    checks.append(AssetCheck("snapshot_total", "Snapshot-Summe vorhanden.", snapshot_total is not None))
    return DocumentAssetVerification(
        document_number=document.document_number,
        ok=all(check.ok for check in checks),
        checks=tuple(checks),
    )


def list_customer_folder_updates(session: Session, customer_id: int) -> list[Path]:
    customer = _customer_or_error(session, customer_id)
    folder = Path(customer.folder_path)
    if not folder.is_dir():
        return []
    known_paths = {
        Path(path).resolve()
        for path in session.scalars(select(Document.excel_path).where(Document.customer_id == customer.id))
        if path
    }
    updates = []
    for path in folder.iterdir():
        if not path.is_file() or path.suffix.lower() != ".xlsx" or path.name.startswith("~$"):
            continue
        if path.resolve() not in known_paths:
            updates.append(path)
    return sorted(updates, key=lambda item: item.name.lower())


def list_customer_folder_excel_previews(session: Session, customer_id: int) -> list[CustomerFolderExcelPreview]:
    return tuple(_preview_customer_excel(path) for path in list_customer_folder_updates(session, customer_id))


def payment_hint_preview(session: Session, customer_id: int, document_date: str) -> PaymentHint:
    customer = _customer_or_error(session, customer_id)
    if _is_bank_transfer(customer.payment_method):
        due_date = _due_date(document_date)
        if due_date is None:
            return PaymentHint(None, "Bitte ueberweisen Sie den Rechnungsbetrag.")
        return PaymentHint(due_date, f"Bitte ueberweisen Sie den Rechnungsbetrag bis zum {due_date}.")
    if _is_sepa(customer.payment_method):
        return PaymentHint(None, "Rechnungsbetrag wird per Sepa Basis Lastschrift Mandat eingezogen.")
    return PaymentHint(None, None)


def list_month_end_checks(session: Session, month: str, today: str) -> list[MonthEndCheck]:
    checks: list[MonthEndCheck] = []
    for item in session.scalars(select(OpenItem).where(OpenItem.status == "offen").order_by(OpenItem.due_date)):
        if (item.document_date or "").startswith(month) and item.due_date and item.due_date < today:
            checks.append(MonthEndCheck("overdue_open_item", f"{item.document_number} ist seit {item.due_date} faellig."))
    documents = list(
        session.scalars(
            select(Document)
            .where(Document.delivery_date.like(f"{month}%"))
            .where(Document.number_released == False)  # noqa: E712
            .order_by(Document.document_number)
        )
    )
    for document in documents:
        verification = verify_document_assets(session, document.id)
        for check in verification.checks:
            if not check.ok:
                checks.append(MonthEndCheck(check.code, f"{document.document_number}: {check.message}", "error"))
        if document.document_type == "Rechnung" and not document.datev_export_path:
            checks.append(MonthEndCheck("missing_datev_export", f"{document.document_number}: DATEV-PDF fehlt."))
    return checks


def suggest_typical_order_lines(session: Session, customer_id: int, recent_order_count: int = 3) -> list[TypicalOrderLine]:
    orders = list(
        session.scalars(
            select(Order)
            .options(selectinload(Order.lines))
            .where(Order.customer_id == customer_id)
            .where(Order.status != "archiviert")
            .order_by(Order.delivery_date.desc(), Order.id.desc())
            .limit(recent_order_count)
        )
    )
    quantities: dict[tuple[int | None, str], list[int]] = {}
    for order in orders:
        for line in order.lines:
            quantities.setdefault((line.product_id, line.product_name), []).append(line.quantity)
    suggestions = [
        TypicalOrderLine(
            product_id,
            name,
            round(sum(values) / len(values)),
            len(values),
            values[0],
            min(values),
            max(values),
            _frequency_label(len(values), max(len(orders), 1)),
        )
        for (product_id, name), values in quantities.items()
    ]
    return sorted(suggestions, key=lambda item: -item.seen_in_orders)


def list_document_sequence_issues(session: Session) -> list[DocumentSequenceIssue]:
    issues: list[DocumentSequenceIssue] = []
    orders = list(
        session.scalars(
            select(Order)
            .options(selectinload(Order.lines))
            .where(Order.status != "archiviert")
            .order_by(Order.delivery_date, Order.order_number)
        )
    )
    for order in orders:
        documents = list(
            session.scalars(
                select(Document)
                .where(Document.order_id == order.id)
                .where(Document.number_released == False)  # noqa: E712
                .order_by(Document.document_type, Document.document_number)
            )
        )
        delivery_notes = [document for document in documents if document.document_type == "Lieferschein"]
        invoices = [document for document in documents if document.document_type == "Rechnung"]
        if delivery_notes and not invoices:
            issues.append(
                DocumentSequenceIssue(
                    "missing_invoice_after_delivery_note",
                    order.order_number,
                    f"Zu Bestellung {order.order_number} gibt es Lieferschein, aber noch keine Rechnung.",
                )
            )
        if len(delivery_notes) > 1:
            issues.append(
                DocumentSequenceIssue(
                    "duplicate_delivery_note_for_order",
                    order.order_number,
                    f"Zu Bestellung {order.order_number} gibt es mehrere Lieferscheine.",
                )
            )
        if len(invoices) > 1:
            issues.append(
                DocumentSequenceIssue(
                    "duplicate_invoice_for_order",
                    order.order_number,
                    f"Zu Bestellung {order.order_number} gibt es mehrere Rechnungen.",
                )
            )
        for document in documents:
            verification = verify_document_assets(session, document.id, expected_assets={"excel", "pdf"})
            if not verification.ok:
                issues.append(
                    DocumentSequenceIssue(
                        "missing_document_asset",
                        order.order_number,
                        f"{document.document_number}: Excel oder PDF fehlt.",
                        "error",
                    )
                )
    return issues


def list_customer_price_deviations(session: Session, customer_id: int | None = None) -> list[CustomerPriceDeviation]:
    query = (
        select(CustomerAssortmentItem)
        .options(selectinload(CustomerAssortmentItem.product))
        .where(CustomerAssortmentItem.is_active == True)  # noqa: E712
        .where(CustomerAssortmentItem.product_id.is_not(None))
    )
    if customer_id is not None:
        query = query.where(CustomerAssortmentItem.customer_id == customer_id)
    deviations: list[CustomerPriceDeviation] = []
    for item in session.scalars(query.order_by(CustomerAssortmentItem.source_product_name)):
        product = item.product
        if product is None or item.last_unit_price_cents == product.standard_price_cents:
            continue
        deviations.append(
            CustomerPriceDeviation(
                item.customer_id,
                product.id,
                product.name,
                item.last_unit_price_cents,
                product.standard_price_cents,
                item.price_decision,
            )
        )
    return deviations


def list_work_cockpit_items(session: Session) -> list[WorkCockpitItem]:
    items: list[WorkCockpitItem] = []
    for issue in session.scalars(select(OnboardingIssue).where(OnboardingIssue.status != "erledigt").order_by(OnboardingIssue.customer_name)):
        items.append(WorkCockpitItem("open_checklist_issue", "Offener Pruefpunkt", f"{issue.customer_name}: {issue.message}"))
    for customer in session.scalars(select(Customer).where(Customer.is_active == True).order_by(Customer.name)):  # noqa: E712
        if not Path(customer.folder_path).is_dir():
            items.append(WorkCockpitItem("missing_customer_folder", "Kundenordner fehlt", customer.name))
        if not customer.contact_email:
            items.append(WorkCockpitItem("missing_contact_email", "Kontakt-Mail fehlt", customer.name, "info"))
    for assortment_item in session.scalars(
        select(CustomerAssortmentItem)
        .options(selectinload(CustomerAssortmentItem.customer))
        .where(CustomerAssortmentItem.is_active == True)  # noqa: E712
        .where(CustomerAssortmentItem.product_id.is_(None))
        .order_by(CustomerAssortmentItem.source_product_name)
    ):
        items.append(
            WorkCockpitItem(
                "unresolved_customer_article",
                "Artikel nicht zugeordnet",
                f"{assortment_item.customer.name}: {assortment_item.source_product_name}",
            )
        )
    for issue in list_document_sequence_issues(session):
        items.append(WorkCockpitItem(issue.code, "Belegfolge pruefen", issue.message, issue.severity))
    for deviation in list_customer_price_deviations(session):
        items.append(
            WorkCockpitItem(
                "customer_price_deviation",
                "Kundenpreis pruefen",
                f"{deviation.product_name}: Kundenpreis weicht vom Artikelstamm ab.",
            )
        )
    return items


def _preview_customer_excel(path: Path) -> CustomerFolderExcelPreview:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
    except Exception:
        return CustomerFolderExcelPreview(path, _document_number_from_filename(path), None, 0)
    document_number = _document_number_from_cells(values) or _document_number_from_filename(path)
    document_date = _first_date_from_cells(values)
    filled_quantity_rows = _filled_quantity_rows(values)
    return CustomerFolderExcelPreview(path, document_number, document_date, filled_quantity_rows)


def _document_number_from_cells(rows: list[tuple]) -> str | None:
    for row in rows:
        for index, value in enumerate(row):
            text = str(value).strip() if value is not None else ""
            if text.lower().replace(" ", "") in {"re.nr.", "renr.", "renr", "lsnr.", "lsnr"}:
                for neighbor in row[index + 1 : index + 4]:
                    if neighbor not in (None, ""):
                        return str(neighbor).strip()
            match = re.search(r"\b(?:RE|LS)[-\s]?\d+\b", text, flags=re.IGNORECASE)
            if match:
                return match.group(0).replace(" ", "-").upper()
    return None


def _document_number_from_filename(path: Path) -> str | None:
    match = re.search(r"\b(?:RE|LS)[-_ ]?\d+\b", path.stem, flags=re.IGNORECASE)
    if not match:
        return None
    return match.group(0).replace("_", "-").replace(" ", "-").upper()


def _first_date_from_cells(rows: list[tuple]) -> str | None:
    for row in rows:
        for value in row:
            if isinstance(value, datetime):
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
            if isinstance(value, str):
                text = value.strip()
                for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
                    try:
                        return datetime.strptime(text, fmt).date().isoformat()
                    except ValueError:
                        pass
    return None


def _filled_quantity_rows(rows: list[tuple]) -> int:
    count = 0
    for row in rows:
        has_name = any(isinstance(value, str) and value.strip() for value in row)
        has_positive_quantity = any(isinstance(value, (int, float)) and value > 0 and float(value).is_integer() for value in row[:4])
        if has_name and has_positive_quantity:
            count += 1
    return count


def _frequency_label(seen_in_orders: int, order_count: int) -> str:
    ratio = seen_in_orders / order_count
    if ratio >= 0.75:
        return "meistens bestellt"
    if ratio >= 0.5:
        return "manchmal bestellt"
    return "selten bestellt"


def _customer_or_error(session: Session, customer_id: int) -> Customer:
    customer = session.get(Customer, customer_id)
    if customer is None:
        raise ValueError("Kunde wurde nicht gefunden.")
    return customer


def _looks_similar(candidate: str, existing: str) -> bool:
    candidate_compact = _compact(candidate)
    existing_compact = _compact(existing)
    return (
        candidate_compact in existing_compact
        or existing_compact in candidate_compact
        or SequenceMatcher(None, candidate_compact, existing_compact).ratio() >= 0.78
    )


def _compact(value: str) -> str:
    return "".join(character for character in value.casefold() if character.isalnum())


def _path_check(code: str, ok_message: str, raw_path: str | None) -> AssetCheck:
    if raw_path and Path(raw_path).exists():
        return AssetCheck(code.replace("missing_", ""), ok_message, True)
    return AssetCheck(code, ok_message, False)


def _is_bank_transfer(payment_method: str | None) -> bool:
    normalized = (payment_method or "").lower().replace("ü", "ue")
    return "ueberweisung" in normalized


def _is_sepa(payment_method: str | None) -> bool:
    return "sepa" in (payment_method or "").lower() or "lastschrift" in (payment_method or "").lower()


def _due_date(document_date: str) -> str | None:
    try:
        return (date.fromisoformat(document_date) + timedelta(days=7)).isoformat()
    except ValueError:
        return None
