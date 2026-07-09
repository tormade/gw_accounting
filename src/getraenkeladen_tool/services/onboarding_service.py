from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Customer, CustomerAssortmentItem, OnboardingIssue, Product, ProductAlias
from ..schemas import CustomerCreate
from .customer_service import create_customer, update_customer
from .master_data_import_service import _clean_text, _date_to_iso, _sheet_rows


ITEM_FIRST_ROW = 13
ITEM_LAST_ROW = 30
RETURN_FIRST_ROW = 34
RETURN_LAST_ROW = 39


@dataclass(slots=True)
class OnboardingLine:
    name: str
    quantity: int
    deposit_cents: int
    unit_price_cents: int
    total_cents: int


@dataclass(slots=True)
class OnboardingDepositReturn:
    quantity: int
    deposit_cents: int
    total_cents: int


@dataclass(slots=True)
class CustomerWorkbookSnapshot:
    source_file: str
    customer_name: str
    address: str | None
    contact_email: str | None
    phone: str | None
    payment_method: str | None
    document_number: str | None
    document_date: str | None
    due_date: str | None
    delivery_comment: str | None
    footer_text: str | None
    lines: list[OnboardingLine] = field(default_factory=list)
    deposit_returns: list[OnboardingDepositReturn] = field(default_factory=list)
    quantity_total: int = 0
    delivery_fee_cents: int = 0
    delivery_total_cents: int = 0
    deposit_return_total_cents: int = 0
    net_total_cents: int = 0
    tax_total_cents: int = 0
    gross_total_cents: int = 0


@dataclass(slots=True)
class CustomerListSnapshot:
    source_file: str
    source_row: int
    name: str
    address: str | None
    contact_email: str | None
    invoice_email: str | None
    phone: str | None
    next_contact_date: str | None
    rhythm: str | None
    priority: str | None
    delivery_notes: str | None
    opening_hours: str | None
    internal_notes: str | None


@dataclass(slots=True)
class OnboardingReport:
    customers_read: int = 0
    assortment_lines: int = 0
    conflicts: int = 0
    unreadable_files: int = 0


@dataclass(slots=True)
class OnboardingResult:
    customer: Customer
    snapshot: CustomerWorkbookSnapshot
    report: OnboardingReport
    issues: list[OnboardingIssue]


@dataclass(slots=True)
class SkippedOnboardingFile:
    path: Path
    message: str


@dataclass(slots=True)
class FolderOnboardingResult:
    report: OnboardingReport
    results: list[OnboardingResult] = field(default_factory=list)
    skipped_files: list[SkippedOnboardingFile] = field(default_factory=list)


@dataclass(frozen=True, slots=True)
class _AnalyzedWorkbook:
    path: Path
    snapshot: CustomerWorkbookSnapshot


def analyze_customer_workbook(path: Path) -> CustomerWorkbookSnapshot:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    text_cells = _text_cells(sheet, max_row=12)
    address_values = [_clean_text(sheet.cell(row=row, column=1).value) for row in range(1, 6)]
    name, address = _name_and_address_from_header(address_values)
    document_date = _find_first_date(sheet)
    document_number = _find_document_number(sheet)
    contact_email = _first_email(text_cells)
    phone = _first_phone(text_cells, exclude={contact_email, name, address})
    footer_text = _footer_text(sheet)
    payment_method = _payment_method_from_footer(footer_text)
    due_date = _due_date_from_footer(sheet, payment_method)
    delivery_comment = _delivery_comment(sheet, {name, address, contact_email, phone})

    lines = _item_lines(sheet)
    deposit_returns = _deposit_returns(sheet)
    delivery_fee_cents = _money_to_cents(sheet["E31"].value)
    totals = _document_totals(sheet)

    return CustomerWorkbookSnapshot(
        source_file=str(path),
        customer_name=name or path.stem,
        address=address,
        contact_email=contact_email,
        phone=phone,
        payment_method=payment_method,
        document_number=document_number,
        document_date=document_date,
        due_date=due_date,
        delivery_comment=delivery_comment,
        footer_text=footer_text,
        lines=lines,
        deposit_returns=deposit_returns,
        quantity_total=totals["quantity_total"],
        delivery_fee_cents=delivery_fee_cents,
        delivery_total_cents=totals["delivery_total_cents"],
        deposit_return_total_cents=totals["deposit_return_total_cents"],
        net_total_cents=totals["net_total_cents"],
        tax_total_cents=totals["tax_total_cents"],
        gross_total_cents=totals["gross_total_cents"],
    )


def onboard_customer_from_sources(
    session: Session,
    customer_name: str,
    customer_list_path: Path,
    workbook_path: Path,
) -> OnboardingResult:
    folder_snapshot = analyze_customer_workbook(workbook_path)
    list_snapshot = read_customer_list_snapshot(customer_list_path, customer_name)
    issues = _build_conflicts(folder_snapshot, list_snapshot)

    payload = CustomerCreate(
        name=list_snapshot.name,
        folder_path=f"Kunden/{list_snapshot.name}",
        address=list_snapshot.address,
        phone=folder_snapshot.phone or list_snapshot.phone,
        contact_email=folder_snapshot.contact_email or list_snapshot.invoice_email or list_snapshot.contact_email,
        payment_method=folder_snapshot.payment_method,
        next_contact_date=list_snapshot.next_contact_date,
        delivery_notes=_join_notes(list_snapshot.delivery_notes, folder_snapshot.delivery_comment),
        opening_hours=list_snapshot.opening_hours,
        internal_notes=list_snapshot.internal_notes,
        source_file=folder_snapshot.source_file,
        source_row=list_snapshot.source_row,
    )
    customer = session.scalar(select(Customer).where(Customer.name == list_snapshot.name))
    if customer is None:
        customer = create_customer(session, payload)
    else:
        customer = update_customer(session, customer.id, payload)

    assortment_issues = _sync_customer_assortment(session, customer, folder_snapshot)
    issues.extend(assortment_issues)

    persisted_issues = []
    for issue in issues:
        session.add(issue)
        persisted_issues.append(issue)
    session.commit()
    for issue in persisted_issues:
        session.refresh(issue)

    return OnboardingResult(
        customer=customer,
        snapshot=folder_snapshot,
        report=OnboardingReport(
            customers_read=1,
            assortment_lines=len(folder_snapshot.lines),
            conflicts=len(persisted_issues),
            unreadable_files=0,
        ),
        issues=persisted_issues,
    )


def _sync_customer_assortment(
    session: Session,
    customer: Customer,
    snapshot: CustomerWorkbookSnapshot,
) -> list[OnboardingIssue]:
    existing_items = {
        _normalize_compare(item.source_product_name): item
        for item in session.scalars(
            select(CustomerAssortmentItem).where(CustomerAssortmentItem.customer_id == customer.id)
        )
    }
    issues = []
    product_index = _product_index(session)
    for sort_order, line in enumerate(snapshot.lines, start=1):
        product = product_index.get(_normalize_product_name(line.name))
        normalized_source = _normalize_compare(line.name)
        item = existing_items.get(normalized_source)
        if item is None:
            item = CustomerAssortmentItem(
                customer_id=customer.id,
                source_product_name=line.name,
                sort_order=sort_order,
                source_file=snapshot.source_file,
            )
            session.add(item)
        item.product_id = product.id if product is not None else None
        item.last_quantity = line.quantity
        item.last_unit_price_cents = line.unit_price_cents
        item.last_deposit_cents = line.deposit_cents
        item.price_decision = item.price_decision or "offen"
        item.sort_order = sort_order
        item.is_active = True
        item.source_file = snapshot.source_file
        if product is None:
            issues.append(
                OnboardingIssue(
                    customer_name=customer.name,
                    source_file=snapshot.source_file,
                    issue_type="product_match",
                    field_name="product",
                    list_value=None,
                    folder_value=line.name,
                    message=f"Artikel konnte nicht sicher zugeordnet werden: {line.name}",
                    status="offen",
                    created_at=datetime.now().isoformat(timespec="seconds"),
                )
            )
        else:
            _ensure_product_alias(session, product, line.name, snapshot.source_file)
            if _prices_differ(product, line):
                item.price_decision = "offen"
                issues.append(
                    OnboardingIssue(
                        customer_name=customer.name,
                        source_file=snapshot.source_file,
                        issue_type="price_mismatch",
                        field_name="price",
                        list_value=(
                            f"{product.name}: zentraler Preis {_format_cents(product.standard_price_cents)}, "
                            f"Pfand {_format_cents(product.default_deposit_cents)}"
                        ),
                        folder_value=(
                            f"{line.name}: Preis aus Excel {_format_cents(line.unit_price_cents)}, "
                            f"Pfand {_format_cents(line.deposit_cents)}"
                        ),
                        message=(
                            "Preisabweichung gefunden. Beim Erstellen einer Bestellung muss entschieden werden, "
                            "ob der zentral gepflegte Preis oder der Preis aus der Excel-Datei genutzt wird."
                        ),
                        status="offen",
                        created_at=datetime.now().isoformat(timespec="seconds"),
                    )
                )
    return issues


def _prices_differ(product: Product, line: OnboardingLine) -> bool:
    return (
        product.standard_price_cents != line.unit_price_cents
        or product.default_deposit_cents != line.deposit_cents
    )


def _format_cents(cents: int) -> str:
    return f"{cents / 100:.2f} EUR".replace(".", ",")


def _product_index(session: Session) -> dict[str, Product]:
    products = list(session.scalars(select(Product).where(Product.is_active == True)))  # noqa: E712
    aliases = list(session.scalars(select(ProductAlias).where(ProductAlias.product_id.is_not(None))))
    index = {_normalize_product_name(product.name): product for product in products}
    products_by_id = {product.id: product for product in products}
    for alias in aliases:
        product = products_by_id.get(alias.product_id)
        if product is not None:
            index[_normalize_product_name(alias.alias)] = product
    return index


def _ensure_product_alias(session: Session, product: Product, alias: str, source_file: str) -> None:
    existing = session.scalar(
        select(ProductAlias)
        .where(ProductAlias.product_id == product.id)
        .where(ProductAlias.alias == alias)
        .limit(1)
    )
    if existing is not None:
        existing.status = "bestaetigt"
        existing.source_file = source_file
        return
    session.add(
        ProductAlias(
            product_id=product.id,
            alias=alias,
            source_file=source_file,
            status="bestaetigt",
        )
    )


def onboard_customer_workbook_folder(
    session: Session,
    customer_list_path: Path,
    folder_path: Path,
) -> FolderOnboardingResult:
    folder_result = FolderOnboardingResult(report=OnboardingReport())
    workbooks_by_customer: dict[str, list[_AnalyzedWorkbook]] = {}
    for workbook_path in sorted(folder_path.rglob("*.xlsx")):
        if _is_excel_lock_file(workbook_path):
            folder_result.skipped_files.append(
                SkippedOnboardingFile(
                    path=workbook_path,
                    message=f"{workbook_path.name}: Excel-Sperrdatei uebersprungen.",
                )
            )
            continue
        try:
            snapshot = analyze_customer_workbook(workbook_path)
        except Exception as error:
            folder_result.report.unreadable_files += 1
            folder_result.skipped_files.append(
                SkippedOnboardingFile(path=workbook_path, message=f"{workbook_path.name}: {error}")
            )
            continue
        customer_key = _normalize_compare(snapshot.customer_name)
        workbooks_by_customer.setdefault(customer_key, []).append(_AnalyzedWorkbook(workbook_path, snapshot))

    latest_workbooks = _latest_workbook_per_customer(workbooks_by_customer)
    for skipped in _older_workbooks(workbooks_by_customer, latest_workbooks):
        folder_result.skipped_files.append(
            SkippedOnboardingFile(
                path=skipped.path,
                message=f"{skipped.path.name}: aeltere Kundenordnerdatei uebersprungen.",
            )
        )

    for workbook in latest_workbooks:
        try:
            result = onboard_customer_from_sources(
                session,
                customer_name=workbook.snapshot.customer_name,
                customer_list_path=customer_list_path,
                workbook_path=workbook.path,
            )
        except Exception as error:
            folder_result.report.unreadable_files += 1
            folder_result.skipped_files.append(
                SkippedOnboardingFile(path=workbook.path, message=f"{workbook.path.name}: {error}")
            )
            continue
        folder_result.results.append(result)
        folder_result.report.customers_read += result.report.customers_read
        folder_result.report.assortment_lines += result.report.assortment_lines
        folder_result.report.conflicts += result.report.conflicts
        folder_result.report.unreadable_files += result.report.unreadable_files
    return folder_result


def _is_excel_lock_file(path: Path) -> bool:
    return path.name.startswith("~$")


def _latest_workbook_per_customer(workbooks_by_customer: dict[str, list[_AnalyzedWorkbook]]) -> list[_AnalyzedWorkbook]:
    latest = []
    for workbooks in workbooks_by_customer.values():
        latest.append(max(workbooks, key=_workbook_recency_key))
    return sorted(latest, key=lambda workbook: workbook.path.as_posix().lower())


def _older_workbooks(
    workbooks_by_customer: dict[str, list[_AnalyzedWorkbook]],
    latest_workbooks: list[_AnalyzedWorkbook],
) -> list[_AnalyzedWorkbook]:
    latest_paths = {workbook.path for workbook in latest_workbooks}
    skipped = []
    for workbooks in workbooks_by_customer.values():
        skipped.extend(workbook for workbook in workbooks if workbook.path not in latest_paths)
    return sorted(skipped, key=lambda workbook: workbook.path.as_posix().lower())


def _workbook_recency_key(workbook: _AnalyzedWorkbook) -> tuple[str, float, str]:
    document_date = workbook.snapshot.document_date or "0000-00-00"
    try:
        modified_at = workbook.path.stat().st_mtime
    except OSError:
        modified_at = 0
    return (document_date, modified_at, workbook.path.name.lower())


def read_customer_list_snapshot(path: Path, customer_name: str) -> CustomerListSnapshot:
    for row_number, row in _sheet_rows(path):
        name = _clean_text(row.get("Name"))
        if name != customer_name:
            continue
        return CustomerListSnapshot(
            source_file=str(path),
            source_row=row_number,
            name=name,
            address=_address_from_customer_list_row(row),
            contact_email=_clean_text(row.get("e-mail (Kontakt)")),
            invoice_email=_clean_text(row.get("e-mail (Re Versand)")),
            phone=_clean_text(row.get("Festnetz")) or _clean_text(row.get("Festnetz ")) or _clean_text(row.get("Handy")),
            next_contact_date=_date_to_iso(row.get("nächster Kontakt")),
            rhythm=_clean_text(row.get("Rhythmus")),
            priority=_clean_text(row.get("Priorität")),
            delivery_notes=_clean_text(row.get("Bemerkungen")),
            opening_hours=_opening_hours_from_customer_list_row(row),
            internal_notes=_clean_text(row.get("ABO ")),
        )
    raise ValueError(f"Kunde {customer_name} wurde in der Lieferkundenliste nicht gefunden.")


def _text_cells(sheet, max_row: int) -> list[str]:
    values = []
    for row in sheet.iter_rows(min_row=1, max_row=max_row, values_only=True):
        values.extend(text for value in row if (text := _clean_text(value)))
    return values


def _name_and_address_from_header(values: list[str | None]) -> tuple[str | None, str | None]:
    parts = [value for value in values if value and value.lower() != "firma"]
    if not parts:
        return None, None
    name = parts[0]
    address = ", ".join(parts[1:]) if len(parts) > 1 else None
    return name, address


def _find_first_date(sheet) -> str | None:
    for row in sheet.iter_rows(min_row=1, max_row=12, values_only=True):
        for value in row:
            if isinstance(value, datetime):
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
    return None


def _find_document_number(sheet) -> str | None:
    for row in sheet.iter_rows(min_row=1, max_row=12):
        for cell in row:
            text = _clean_text(cell.value)
            if text and "re. nr" in text.lower():
                value = sheet.cell(row=cell.row, column=cell.column + 1).value
                return _clean_text(value)
    return None


def _first_email(values: list[str]) -> str | None:
    for value in values:
        if "@" in value:
            return value
    return None


def _first_phone(values: list[str], exclude: set[str | None]) -> str | None:
    for value in values:
        if value in exclude:
            continue
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", value):
            continue
        if re.fullmatch(r"\d+", value):
            continue
        digits = sum(character.isdigit() for character in value)
        has_phone_separator = any(character in value for character in (" ", "-", "/", ",", "(", ")"))
        if digits >= 6 and has_phone_separator and not _looks_like_address(value):
            return value
    return None


def _looks_like_address(value: str) -> bool:
    lowered = value.lower()
    return bool(re.search(r"\b(str\.?|straße|weg|platz)\b", lowered)) or bool(
        re.fullmatch(r"\d{5}\s+[^\d].*", value)
    )


def _footer_text(sheet) -> str | None:
    footer_values = []
    for row in sheet.iter_rows(min_row=40, max_row=48, values_only=True):
        for value in row:
            text = _clean_text(value)
            if text and (
                "sepa" in text.lower()
                or "überweisen" in text.lower()
                or "ueberweisen" in text.lower()
                or "ware bleibt" in text.lower()
            ):
                footer_values.append(text)
    return " ".join(footer_values) or None


def _payment_method_from_footer(footer_text: str | None) -> str | None:
    lowered = (footer_text or "").lower()
    if "sepa" in lowered or "lastschrift" in lowered:
        return "SEPA"
    if "überweisen" in lowered or "ueberweisen" in lowered:
        return "Überweisung"
    return None


def _due_date_from_footer(sheet, payment_method: str | None) -> str | None:
    if payment_method != "Überweisung":
        return None
    for row in sheet.iter_rows(min_row=40, max_row=48):
        for cell in row:
            text = _clean_text(cell.value)
            if not text:
                continue
            lowered = text.lower()
            if "überweisen" not in lowered and "ueberweisen" not in lowered:
                continue
            for candidate in row:
                if isinstance(candidate.value, datetime):
                    return candidate.value.date().isoformat()
                if isinstance(candidate.value, date):
                    return candidate.value.isoformat()
    return None


def _delivery_comment(sheet, excluded_values: set[str | None]) -> str | None:
    for cell in sheet[1]:
        text = _clean_text(cell.value)
        if text and text not in excluded_values and text.lower() != "firma":
            return text
    return None


def _item_lines(sheet) -> list[OnboardingLine]:
    lines = []
    for row in range(ITEM_FIRST_ROW, ITEM_LAST_ROW + 1):
        name = _clean_text(sheet.cell(row=row, column=2).value)
        if not name:
            continue
        lines.append(
            OnboardingLine(
                name=name,
                quantity=_int_value(sheet.cell(row=row, column=1).value),
                deposit_cents=_money_to_cents(sheet.cell(row=row, column=3).value),
                unit_price_cents=_money_to_cents(sheet.cell(row=row, column=4).value),
                total_cents=_money_to_cents(sheet.cell(row=row, column=5).value),
            )
        )
    return lines


def _deposit_returns(sheet) -> list[OnboardingDepositReturn]:
    returns = []
    for row in range(RETURN_FIRST_ROW, RETURN_LAST_ROW + 1):
        quantity = _int_value(sheet.cell(row=row, column=1).value)
        deposit_cents = _money_to_cents(sheet.cell(row=row, column=3).value)
        total_cents = _money_to_cents(sheet.cell(row=row, column=5).value)
        if quantity == 0 and deposit_cents == 0 and total_cents == 0:
            continue
        returns.append(OnboardingDepositReturn(quantity=quantity, deposit_cents=deposit_cents, total_cents=total_cents))
    return returns


def _document_totals(sheet) -> dict[str, int]:
    totals = {
        "quantity_total": 0,
        "delivery_total_cents": 0,
        "deposit_return_total_cents": 0,
        "net_total_cents": 0,
        "tax_total_cents": 0,
        "gross_total_cents": 0,
    }
    for row in range(25, 45):
        row_values = [sheet.cell(row=row, column=column).value for column in range(1, 7)]
        row_text = " ".join(_clean_text(value) or "" for value in row_values).lower()
        if "lieferwert" in row_text:
            totals["quantity_total"] = _first_int_in_values(row_values)
            totals["delivery_total_cents"] = _last_money_in_values(row_values)
        elif "pfand" in row_text and "rückgabe" in row_text:
            totals["deposit_return_total_cents"] = _last_money_in_values(row_values)
        elif "netto" in row_text:
            totals["net_total_cents"] = _last_money_in_values(row_values)
        elif "mehrwertsteuer" in row_text:
            totals["tax_total_cents"] = _last_money_in_values(row_values)
        elif "brutto" in row_text:
            totals["gross_total_cents"] = _last_money_in_values(row_values)
    return totals


def _first_int_in_values(values: list) -> int:
    for value in values:
        if isinstance(value, int | float) and not isinstance(value, bool):
            return int(value)
    return 0


def _last_money_in_values(values: list) -> int:
    for value in reversed(values):
        cents = _money_to_cents(value)
        if cents != 0:
            return cents
    return 0


def _address_from_customer_list_row(row: dict) -> str | None:
    street = _clean_text(row.get("Straße"))
    house_number = _clean_text(row.get("Hausnummer"))
    postal_code = _clean_text(row.get("PLZ"))
    city = _clean_text(row.get("Ort"))
    street_part = " ".join(part for part in (street, house_number) if part)
    city_part = " ".join(part for part in (postal_code, city) if part)
    return ", ".join(part for part in (street_part, city_part) if part) or None


def _opening_hours_from_customer_list_row(row: dict) -> str | None:
    values = []
    if _clean_text(row.get("Vormittag")):
        values.append("vormittag")
    if _clean_text(row.get("Nachmittag")) or _clean_text(row.get("Nachmittag ")):
        values.append("nachmittag")
    return ", ".join(values) or None


def _build_conflicts(
    folder_snapshot: CustomerWorkbookSnapshot,
    list_snapshot: CustomerListSnapshot,
) -> list[OnboardingIssue]:
    issues = []
    for field_name, list_value, folder_value in (
        ("address", list_snapshot.address, folder_snapshot.address),
        ("phone", list_snapshot.phone, folder_snapshot.phone),
        ("contact_email", list_snapshot.contact_email or list_snapshot.invoice_email, folder_snapshot.contact_email),
    ):
        if not list_value or not folder_value:
            continue
        if _normalize_compare(list_value) == _normalize_compare(folder_value):
            continue
        issues.append(
            OnboardingIssue(
                customer_name=list_snapshot.name,
                source_file=folder_snapshot.source_file,
                issue_type="merge_conflict",
                field_name=field_name,
                list_value=list_value,
                folder_value=folder_value,
                message=f"{field_name} unterscheidet sich zwischen Lieferkundenliste und Kundenordnerdatei.",
                status="offen",
                created_at=datetime.now().isoformat(timespec="seconds"),
            )
        )
    return issues


def _join_notes(*values: str | None) -> str | None:
    notes = [value for value in values if value]
    return " | ".join(notes) if notes else None


def _normalize_compare(value: str) -> str:
    return re.sub(r"[^a-z0-9@]+", "", value.lower())


def _normalize_product_name(value: str) -> str:
    normalized = value.lower()
    normalized = normalized.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    normalized = normalized.replace(".", "")
    normalized = re.sub(r"\s+", " ", normalized)
    normalized = normalized.replace(" x ", "x")
    normalized = normalized.replace(" x", "x").replace("x ", "x")
    return re.sub(r"[^a-z0-9,]+", "", normalized)


def _int_value(value) -> int:
    if value in (None, "", "€"):
        return 0
    return int(float(value))


def _money_to_cents(value) -> int:
    if value in (None, "", "€"):
        return 0
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return 0
    return int((amount * 100).to_integral_value(rounding=ROUND_HALF_UP))
