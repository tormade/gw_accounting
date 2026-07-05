from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from getraenkeladen_tool.services.excel_service import (
    TEMPLATE_PATH,
    build_delivery_note_workbook,
    build_invoice_workbook,
)

INPUT_DIR = Path("/Users/thomasrumel/Documents/Codex/2026-06-20/Input")


def test_winklmeier_template_is_available():
    assert TEMPLATE_PATH.exists()


def test_build_invoice_workbook_writes_customer_excel_file(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "RG-1001.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="RG-1001",
        document_date="2026-06-22",
        line_items=[{"name": "Wasser", "quantity": 10, "unit_price_cents": 1299}],
    )

    workbook = load_workbook(output_path)
    sheet = workbook.active
    assert sheet["A8"].value == "Rechnung"
    assert sheet["F8"].value == "RG-1001"
    assert sheet["C5"].value.date() == date(2026, 6, 22)
    assert sheet["A13"].value == 10
    assert sheet["B13"].value == "Wasser"
    assert sheet["D13"].value == 12.99
    assert sheet["E13"].value == "=(C13+D13)*A13"
    assert sheet["A31"].value == 0
    assert sheet["B31"].value == "Lieferpauschale"
    assert sheet["D31"].value == 3.9
    assert sheet["E31"].value == "=(C31+D31)*A31"
    assert sheet["F32"].value == "=SUM(E13:E31)"
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.forceFullCalc is True
    assert sheet.print_area == "'Tabelle1'!$A$1:$F$48"


def test_build_invoice_workbook_keeps_zero_quantity_lines_with_blank_quantity_cell(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "RG-NULL.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="RG-NULL",
        document_date="2026-06-22",
        line_items=[
            {"name": "Wasser", "quantity": 0, "unit_price_cents": 1299, "deposit_cents": 330},
            {"name": "Spezi", "quantity": 2, "unit_price_cents": 1599, "deposit_cents": 310},
        ],
    )

    formula_sheet = load_workbook(output_path, data_only=False).active
    assert formula_sheet["A13"].value is None
    assert formula_sheet["B13"].value == "Wasser"
    assert formula_sheet["E13"].value == "=(C13+D13)*A13"
    assert formula_sheet["A14"].value == 2
    assert formula_sheet["B14"].value == "Spezi"

    value_sheet = load_workbook(output_path, data_only=True).active
    assert value_sheet["E13"].value == 0
    assert value_sheet["F43"].value == 38.18


def test_build_delivery_order_workbook_writes_customer_excel_file(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "LS-2001.xlsx"

    build_delivery_note_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="LS-2001",
        document_date="2026-06-22",
        line_items=[{"name": "Apfelschorle", "quantity": 4, "unit_price_cents": 1499}],
    )

    workbook = load_workbook(output_path)
    sheet = workbook.active
    assert sheet["A8"].value == "Lieferschein"
    assert sheet["F8"].value == "LS-2001"
    assert sheet["C5"].value.date() == date(2026, 6, 22)
    assert sheet["A13"].value == 4
    assert sheet["B13"].value == "Apfelschorle"


def test_delivery_note_workbook_writes_delivery_comment_footer_and_fee_choice(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Metzgerei Karl" / "LS-2002.xlsx"

    build_delivery_note_workbook(
        output_path=output_path,
        customer_name="Metzgerei Karl",
        document_number="LS-2002",
        document_date="2026-06-22",
        line_items=[{"name": "Frucade Colamix 20x0,5", "quantity": 3, "unit_price_cents": 1048, "deposit_cents": 310}],
        delivery_comment="bis13Uhr und ab 15 Uhr",
        delivery_fee_enabled=True,
    )

    sheet = load_workbook(output_path, data_only=False).active
    assert sheet["D1"].value == "bis13Uhr und ab 15 Uhr"
    assert sheet["A31"].value == 1
    assert sheet["A45"].value == "Die Ware bleibt bis zur vollstaendigen Bezahlung Eigentum von Getraenke Winklmeier."
    assert sheet["B48"].value == "Vielen Dank fuer Ihren Einkauf"


def test_invoice_workbook_writes_footer_and_optional_delivery_fee(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Metzgerei Karl" / "RG-2002.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Metzgerei Karl",
        document_number="RG-2002",
        document_date="2026-06-22",
        line_items=[{"name": "Frucade Colamix 20x0,5", "quantity": 3, "unit_price_cents": 1048, "deposit_cents": 310}],
        invoice_footer_text="Rechnungsbetrag wird per Sepa Basis Lastschrift Mandat eingezogen.",
        delivery_fee_enabled=True,
    )

    formula_sheet = load_workbook(output_path, data_only=False).active
    assert formula_sheet["A31"].value == 1
    assert formula_sheet["B31"].value == "Lieferpauschale"
    assert formula_sheet["D31"].value == 3.9
    assert formula_sheet["A45"].value == "Rechnungsbetrag wird per Sepa Basis Lastschrift Mandat eingezogen."
    assert formula_sheet["B48"].value == "Vielen Dank fuer Ihren Einkauf"

    value_sheet = load_workbook(output_path, data_only=True).active
    assert value_sheet["E31"].value == 3.9
    assert value_sheet["F43"].value == 44.64


def test_build_invoice_workbook_refreshes_sum_formulas_for_all_item_rows(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "RG-1002.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="RG-1002",
        document_date="2026-06-22",
        line_items=[
            {"name": "Wasser", "quantity": 2, "unit_price_cents": 1299, "deposit_cents": 330},
            {"name": "Spezi", "quantity": 3, "unit_price_cents": 1599, "deposit_cents": 330},
        ],
    )

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook.active

    assert sheet["E13"].value == "=(C13+D13)*A13"
    assert sheet["E14"].value == "=(C14+D14)*A14"
    assert sheet["E30"].value == "=(C30+D30)*A30"
    assert sheet["F32"].value == "=SUM(E13:E31)"
    assert sheet["A32"].value == "=SUM(A13:A30)"
    assert sheet["F40"].value == "=SUM(E34:E39)"
    assert sheet["F41"].value == "=ROUND(F43/1.19,2)"
    assert sheet["F42"].value == "=F43-F41"
    assert sheet["F43"].value == "=F32+F40"


def test_build_invoice_workbook_stores_cached_totals_for_preview_and_data_only_reads(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "RG-1003.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="RG-1003",
        document_date="2026-06-22",
        line_items=[
            {"name": "Wasser", "quantity": 1, "unit_price_cents": 1030, "deposit_cents": 480},
            {"name": "Apfelschorle", "quantity": 1, "unit_price_cents": 1680, "deposit_cents": 480},
            {"name": "Black Forest Still", "quantity": 3, "unit_price_cents": 1180, "deposit_cents": 330},
        ],
    )

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook.active

    assert sheet["E13"].value == 15.1
    assert sheet["E14"].value == 21.6
    assert sheet["E15"].value == 45.3
    assert sheet["A32"].value == 5
    assert sheet["F32"].value == 82
    assert sheet["F40"].value == 0
    assert sheet["F41"].value == 68.91
    assert sheet["F42"].value == 13.09
    assert sheet["F43"].value == 82


def test_build_invoice_workbook_writes_deposit_returns_into_return_block(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Cafe Nord" / "RG-1004.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Cafe Nord",
        document_number="RG-1004",
        document_date="2026-06-22",
        line_items=[{"name": "Wasser", "quantity": 1, "unit_price_cents": 1030, "deposit_cents": 480}],
        deposit_returns=[{"name": "Leergut Kiste 4,80", "quantity": 1, "deposit_cents": 480}],
    )

    formula_workbook = load_workbook(output_path, data_only=False)
    formula_sheet = formula_workbook.active
    assert formula_sheet["A34"].value == 1
    assert formula_sheet["B34"].value == "Leergut Kiste 4,80"
    assert formula_sheet["C34"].value == -4.8
    assert formula_sheet["E34"].value == "=(C34+D34)*A34"
    assert formula_sheet["F40"].value == "=SUM(E34:E39)"
    assert formula_sheet["F43"].value == "=F32+F40"

    value_workbook = load_workbook(output_path, data_only=True)
    value_sheet = value_workbook.active
    assert value_sheet["E34"].value == -4.8
    assert value_sheet["F32"].value == 15.1
    assert value_sheet["F40"].value == -4.8
    assert value_sheet["F43"].value == 10.3


def test_build_invoice_workbook_continues_last_customer_excel_without_removing_article_rows(tmp_path: Path):
    output_path = tmp_path / "Kunden" / "Metzgerei Karl" / "RE-NEU.xlsx"

    build_invoice_workbook(
        output_path=output_path,
        customer_name="Metzgerei Karl",
        document_number="RE-NEU",
        document_date="2026-07-05",
        line_items=[
            {"name": "Frucade Colamix 20x0,5", "quantity": 5, "unit_price_cents": 1048, "deposit_cents": 310},
            {"name": "Brunnthaler Spritzig 20x0,5", "quantity": 2, "unit_price_cents": 748, "deposit_cents": 310},
        ],
        deposit_returns=[
            {"name": "Pfand 3,10 EUR", "quantity": 6, "deposit_cents": 310},
            {"name": "Pfand 4,50 EUR", "quantity": 1, "deposit_cents": 450},
            {"name": "Pfand 1,50 EUR", "quantity": 2, "deposit_cents": 150},
        ],
        source_workbook_path=INPUT_DIR / "_ RE 0525 Metzgerei Karl .xlsx",
    )

    formula_sheet = load_workbook(output_path, data_only=False).active
    assert formula_sheet["A8"].value == "Rechnung"
    assert formula_sheet["F8"].value == "RE-NEU"
    assert formula_sheet["C5"].value.date() == date(2026, 7, 5)
    assert formula_sheet["B13"].value == "Labertaler Apfelschorle 20x0,5"
    assert formula_sheet["A13"].value is None
    assert formula_sheet["B14"].value == "Frucade Colamix 20x0,5"
    assert formula_sheet["A14"].value == 5
    assert formula_sheet["B20"].value == "Brunnthaler Spritzig 20x0,5"
    assert formula_sheet["A20"].value == 2
    assert formula_sheet["B27"].value == "Hofbrauhaus Moy Hell 20x0,5"
    assert formula_sheet["A27"].value is None
    assert formula_sheet["E14"].value == "=(C14+D14)*A14"
    assert formula_sheet["A34"].value == 6
    assert formula_sheet["C34"].value == -3.1
    assert formula_sheet["A35"].value == 1
    assert formula_sheet["C35"].value == -4.5
    assert formula_sheet["A36"].value == 2
    assert formula_sheet["C36"].value == -1.5

    value_sheet = load_workbook(output_path, data_only=True).active
    assert value_sheet["A32"].value == 7
    assert value_sheet["F40"].value == -26.1
