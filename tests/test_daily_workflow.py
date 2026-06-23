from pathlib import Path

from openpyxl import load_workbook

from getraenkeladen_tool.models import OpenItem
from getraenkeladen_tool.schemas import OrderCreate, OrderLineCreate
from getraenkeladen_tool.services.customer_assortment_service import list_customer_assortment
from getraenkeladen_tool.services.customer_folder_service import get_customer_folder_snapshot
from getraenkeladen_tool.services.master_data_import_service import import_master_data_from_folder
from getraenkeladen_tool.services.onboarding_service import onboard_customer_from_sources
from getraenkeladen_tool.services.order_service import create_order, create_order_delivery_order, create_order_invoice


INPUT_DIR = Path("/Users/thomasrumel/Documents/Codex/2026-06-20/Input")


def test_daily_customer_folder_workflow_updates_old_excel_with_central_prices_and_exports_documents(session, tmp_path):
    import_master_data_from_folder(session, INPUT_DIR)
    onboarding = onboard_customer_from_sources(
        session,
        customer_name="Metzgerei Karl",
        customer_list_path=INPUT_DIR / "Lieferkunden Liste.xlsx",
        workbook_path=INPUT_DIR / "_ RE 0525 Metzgerei Karl .xlsx",
    )
    onboarding.customer.folder_path = str(tmp_path / "Kunden" / "Metzgerei Karl")
    session.commit()

    assortment = list_customer_assortment(session, onboarding.customer.id)
    frucade = next(row for row in assortment if row.source_product_name == "Frucade Colamix 20x0,5")
    assert frucade.excel_price_cents == 1048
    assert frucade.current_price_cents == 1190
    assert frucade.price_differs_from_central is True

    order = create_order(
        session,
        OrderCreate(
            order_number="BEST-TEST-1",
            customer_id=onboarding.customer.id,
            order_date="2026-06-23",
            delivery_date="2026-06-24",
            delivery_slot="vormittag",
            lines=[
                OrderLineCreate(
                    product_id=frucade.product_id,
                    quantity=frucade.last_quantity,
                    unit_price_cents=frucade.current_price_cents,
                    deposit_cents=frucade.current_deposit_cents,
                )
            ],
        ),
    )

    delivery_note = create_order_delivery_order(
        session,
        order.id,
        "LS-WORKFLOW-1",
        delivery_fee_enabled=True,
        delivery_comment="bis 13 Uhr",
    )
    invoice = create_order_invoice(
        session,
        order.id,
        "RG-WORKFLOW-1",
        delivery_fee_enabled=True,
        datev_upload_dir=tmp_path / "DATEV",
    )

    assert Path(delivery_note.excel_path).exists()
    assert Path(delivery_note.pdf_path).read_bytes().startswith(b"%PDF-")
    assert Path(invoice.excel_path).exists()
    assert Path(invoice.pdf_path).read_bytes().startswith(b"%PDF-")
    assert invoice.datev_export_path is not None
    snapshot = get_customer_folder_snapshot(session, onboarding.customer.id)
    generated_paths = {file.path for file in snapshot.files}
    assert Path(delivery_note.excel_path).parent == Path(onboarding.customer.folder_path)
    assert Path(invoice.excel_path).parent == Path(onboarding.customer.folder_path)
    assert Path(delivery_note.excel_path) in generated_paths
    assert Path(delivery_note.pdf_path) in generated_paths
    assert Path(invoice.excel_path) in generated_paths
    assert Path(invoice.pdf_path) in generated_paths

    workbook = load_workbook(invoice.excel_path, data_only=False)
    sheet = workbook.active
    assert sheet["B13"].value == "Frucade Colamix 20x0,5"
    assert sheet["A13"].value == 3
    assert sheet["D13"].value == 11.9
    assert sheet["E13"].value == "=(C13+D13)*A13"
    assert sheet["A31"].value == 1
    assert sheet["B31"].value == "Lieferpauschale"
    assert sheet["F32"].value == "=SUM(E13:E31)"

    cached_sheet = load_workbook(invoice.excel_path, data_only=True).active
    assert cached_sheet["F32"].value == 48.9
    assert cached_sheet["F43"].value == 48.9
    assert session.query(OpenItem).one().amount_cents == 4890
