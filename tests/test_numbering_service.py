from pathlib import Path

from getraenkeladen_tool.schemas import CustomerCreate, DocumentCreate, DocumentLineItem, OrderCreate, OrderLineCreate, ProductCreate
from getraenkeladen_tool.services.customer_service import create_customer
from getraenkeladen_tool.services.document_service import create_document
from getraenkeladen_tool.models import OpenItem
from getraenkeladen_tool.services.numbering_service import (
    NumberSuggestions,
    check_number_conflict,
    load_number_sequences_from_workbook,
    list_number_sequence_statuses,
    next_document_number,
    next_number_for_prefix,
    release_number,
    reset_number_sequences_to_defaults,
    set_next_number,
    suggest_next_numbers,
    write_number_sequences_to_workbook,
)
from getraenkeladen_tool.services.order_service import create_order
from getraenkeladen_tool.services.product_service import create_product
from openpyxl import load_workbook


def test_next_number_for_prefix_uses_default_when_no_numbers_exist():
    assert next_number_for_prefix([], prefix="RG", start=3001) == "RG-3001"


def test_next_number_for_prefix_continues_after_highest_matching_number():
    numbers = ["RG-3001", "RG-4500", "LS-9999", "RG-alt"]

    assert next_number_for_prefix(numbers, prefix="RG", start=3001) == "RG-4501"


def test_suggest_next_numbers_uses_database_orders_and_documents(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Nord", folder_path=str(tmp_path / "Cafe Nord")))
    product = create_product(session, ProductCreate(name="Wasser", unit="Kiste", standard_price_cents=1299))
    create_order(
        session,
        OrderCreate(
            order_number="AUF-2500",
            customer_id=customer.id,
            order_date="2026-06-21",
            delivery_date="2026-06-22",
            lines=[OrderLineCreate(product_id=product.id, quantity=1)],
        ),
    )
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Lieferschein",
            document_number="LS-4100",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-4500",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )

    assert suggest_next_numbers(session) == NumberSuggestions(
        order_number="AUF-2501",
        delivery_note_number="LS-4101",
        invoice_number="RG-4501",
    )


def test_suggest_next_numbers_counts_legacy_delivery_order_documents(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Alt", folder_path=str(tmp_path / "Cafe Alt")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Lieferauftrag",
            document_number="LS-4999",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )

    assert suggest_next_numbers(session).delivery_note_number == "LS-5000"


def test_next_document_number_returns_default_for_document_type_without_existing_documents(session):
    assert next_document_number(session, document_type="Rechnung", prefix="RG", start=3001) == "RG-3001"


def test_set_next_number_changes_future_suggestion_when_it_is_higher_than_existing(session):
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-4100")

    assert suggest_next_numbers(session).invoice_number == "RG-4100"


def test_number_suggestion_can_go_below_higher_existing_documents_when_number_is_free(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Nummer", folder_path=str(tmp_path / "Cafe Nummer")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-4100",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-3001")

    assert suggest_next_numbers(session).invoice_number == "RG-3001"


def test_manual_next_number_can_reset_to_free_lower_number(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Reset", folder_path=str(tmp_path / "Cafe Reset")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-4100",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-3001")

    assert suggest_next_numbers(session).invoice_number == "RG-3001"


def test_manual_next_number_skips_used_numbers_when_counting_forward(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Kollision", folder_path=str(tmp_path / "Cafe Kollision")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-3001",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-3001")

    assert suggest_next_numbers(session).invoice_number == "RG-3002"


def test_number_sequence_status_shows_saved_start_value_even_when_blocked(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Anzeige", folder_path=str(tmp_path / "Cafe Anzeige")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-3001",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-3001")

    statuses = {status.sequence_key: status for status in list_number_sequence_statuses(session)}

    assert statuses["invoice"].next_number == "RG-3001"
    assert suggest_next_numbers(session).invoice_number == "RG-3002"


def test_check_number_conflict_finds_blocking_invoice_document(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Konflikt", folder_path=str(tmp_path / "Cafe Konflikt")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-3001",
            delivery_date="2026-06-22",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )

    conflict = check_number_conflict(session, sequence_key="invoice", value="RG-3001")

    assert conflict is not None
    assert conflict.number == "RG-3001"
    assert conflict.kind == "Rechnung"
    assert conflict.customer_name == "Cafe Konflikt"


def test_release_number_allows_invoice_sequence_to_reuse_released_number(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Freigabe", folder_path=str(tmp_path / "Cafe Freigabe")))
    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-3001",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299)],
        ),
    )
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-3001")

    release_number(session, sequence_key="invoice", value="RG-3001")

    assert check_number_conflict(session, sequence_key="invoice", value="RG-3001") is None
    assert suggest_next_numbers(session).invoice_number == "RG-3001"
    assert session.query(OpenItem).one().status == "freigegeben"


def test_archived_orders_do_not_block_order_number_sequence(session, tmp_path: Path):
    customer = create_customer(session, CustomerCreate(name="Cafe Auftrag", folder_path=str(tmp_path / "Cafe Auftrag")))
    product = create_product(session, ProductCreate(name="Wasser", unit="Menge", standard_price_cents=1299))
    order = create_order(
        session,
        OrderCreate(
            order_number="AUF-1001",
            customer_id=customer.id,
            order_date="2026-06-21",
            delivery_date="2026-06-22",
            lines=[OrderLineCreate(product_id=product.id, quantity=1)],
        ),
    )
    order.status = "archiviert"
    session.commit()
    set_next_number(session, sequence_key="order", prefix="AUF", value="AUF-1001")

    assert suggest_next_numbers(session).order_number == "AUF-1001"


def test_reset_number_sequences_to_defaults_restores_all_start_values(session):
    set_next_number(session, sequence_key="order", prefix="AUF", value="AUF-2222")
    set_next_number(session, sequence_key="delivery_note", prefix="LS", value="LS-4444")
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-5555")

    reset_number_sequences_to_defaults(session)

    statuses = {status.sequence_key: status for status in list_number_sequence_statuses(session)}
    assert statuses["order"].next_number == "AUF-1001"
    assert statuses["delivery_note"].next_number == "LS-3001"
    assert statuses["invoice"].next_number == "RG-3001"


def test_number_sequence_workbook_is_leading_source_after_excel_edit(session, tmp_path: Path):
    workbook_path = tmp_path / "Nummernkreise.xlsx"
    set_next_number(session, sequence_key="invoice", prefix="RG", value="RG-5555")
    write_number_sequences_to_workbook(session, workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Nummernkreise"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=2).value == "invoice":
            sheet.cell(row=row, column=4).value = "RG-3333"
    workbook.save(workbook_path)

    load_number_sequences_from_workbook(session, workbook_path)

    assert list_number_sequence_statuses(session)[2].next_number == "RG-3333"
