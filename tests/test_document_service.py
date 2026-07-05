from pathlib import Path

from openpyxl import load_workbook

from getraenkeladen_tool.models import DocumentLine, OpenItem
from getraenkeladen_tool.schemas import CustomerCreate, DepositReturnCreate, DocumentCreate, DocumentLineItem
from getraenkeladen_tool.services.customer_service import create_customer
from getraenkeladen_tool.services.document_service import create_document, latest_invoice_number

INPUT_DIR = Path("/Users/thomasrumel/Documents/Codex/2026-06-20/Input")


def test_document_outputs_use_shared_core_calculation():
    document_source = Path("src/getraenkeladen_tool/services/document_service.py").read_text(encoding="utf-8")
    excel_source = Path("src/getraenkeladen_tool/services/excel_service.py").read_text(encoding="utf-8")
    pdf_source = Path("src/getraenkeladen_tool/services/pdf_service.py").read_text(encoding="utf-8")

    assert "berechne_beleg" in document_source
    assert "berechne_beleg" in excel_source
    assert "berechne_beleg" in pdf_source


def test_create_invoice_writes_excel_pdf_file_and_open_item(session, tmp_path: Path):
    customer_folder = tmp_path / "Kunden" / "Cafe Nord"
    customer = create_customer(
        session,
        CustomerCreate(
            name="Cafe Nord",
            folder_path=str(customer_folder),
            payment_method="SEPA",
        ),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-1001",
            delivery_date="2026-06-21",
            line_items=[
                DocumentLineItem(
                    name="Wasser 0,7",
                    quantity=10,
                    unit_price_cents=1299,
                    deposit_cents=330,
                )
            ],
        ),
    )

    excel_path = Path(document.excel_path)
    assert excel_path == customer_folder / "2026-06-21_RE_RG-1001_Cafe_Nord.xlsx"
    assert excel_path.exists()
    workbook = load_workbook(excel_path, data_only=False)
    assert workbook.active["C5"].value.strftime("%Y-%m-%d") == "2026-06-21"

    pdf_path = Path(document.pdf_path)
    assert pdf_path == customer_folder / "2026-06-21_RE_RG-1001_Cafe_Nord.pdf"
    assert pdf_path.exists()
    assert pdf_path.read_bytes().startswith(b"%PDF-")

    open_item = session.query(OpenItem).one()
    assert open_item.customer_name == "Cafe Nord"
    assert open_item.document_number == "RG-1001"
    assert open_item.amount_cents == 16290
    assert open_item.payment_method == "SEPA"
    assert open_item.status == "offen"


def test_create_invoice_includes_optional_delivery_fee_and_document_texts(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Metzgerei Karl", folder_path=str(tmp_path / "Kunden" / "Metzgerei Karl")),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-2001",
            delivery_date="2026-06-22",
            line_items=[DocumentLineItem(name="Frucade Colamix 20x0,5", quantity=3, unit_price_cents=1048, deposit_cents=310)],
            delivery_fee_enabled=True,
            footer_text="Rechnungsbetrag wird per Sepa Basis Lastschrift Mandat eingezogen.",
        ),
    )

    assert document.delivery_fee_enabled is True
    assert document.footer_text == "Rechnungsbetrag wird per Sepa Basis Lastschrift Mandat eingezogen."
    assert session.query(OpenItem).one().amount_cents == 4464

    sheet = load_workbook(document.excel_path, data_only=True).active
    assert sheet["A31"].value == 1
    assert sheet["E31"].value == 3.9
    assert sheet["F43"].value == 44.64


def test_create_invoice_persists_reproducible_document_line_snapshot(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Snapshot Kunde", folder_path=str(tmp_path / "Kunden" / "Snapshot Kunde")),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-SNAP-1",
            delivery_date="2026-06-22",
            line_items=[
                DocumentLineItem(name="Frucade Colamix 20x0,5", quantity=3, unit_price_cents=1048, deposit_cents=310),
                DocumentLineItem(name="Adelholzener Classic 12x0,7", quantity=5, unit_price_cents=1128, deposit_cents=330),
            ],
            deposit_returns=[DepositReturnCreate(name="Leergut Kiste 3,10", quantity=3, deposit_cents=310)],
            delivery_fee_enabled=True,
        ),
    )

    snapshot_lines = (
        session.query(DocumentLine)
        .filter(DocumentLine.document_id == document.id)
        .order_by(DocumentLine.sort_order)
        .all()
    )

    assert [
        (
            line.line_type,
            line.name,
            line.quantity,
            line.unit_price_cents,
            line.deposit_cents,
            line.total_cents,
        )
        for line in snapshot_lines
    ] == [
        ("position", "Frucade Colamix 20x0,5", 3, 1048, 310, 4074),
        ("position", "Adelholzener Classic 12x0,7", 5, 1128, 330, 7290),
        ("deposit_return", "Leergut Kiste 3,10", 3, 0, 310, -930),
        ("delivery_fee", "Lieferpauschale", 1, 390, 0, 390),
        ("summary", "Brutto", 1, 10824, 0, 10824),
    ]

    open_item = session.query(OpenItem).one()
    assert open_item.amount_cents == 10824


def test_create_invoice_supports_multiple_line_items(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(
            name="Mehrpositionen Kunde",
            folder_path=str(tmp_path / "Kunden" / "Mehrpositionen Kunde"),
            payment_method="Ueberweisung",
        ),
    )

    create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-1002",
            delivery_date="2026-06-21",
            line_items=[
                DocumentLineItem(
                    name="Wasser 0,7",
                    quantity=2,
                    unit_price_cents=1299,
                    deposit_cents=330,
                ),
                DocumentLineItem(
                    name="Apfelschorle",
                    quantity=3,
                    unit_price_cents=1499,
                    deposit_cents=330,
                ),
            ],
        ),
    )

    open_item = session.query(OpenItem).one()
    assert open_item.amount_cents == 8745


def test_create_invoice_reduces_open_item_by_deposit_returns(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(
            name="Pfandkunde",
            folder_path=str(tmp_path / "Kunden" / "Pfandkunde"),
            payment_method="bar",
        ),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-1002-P",
            delivery_date="2026-06-21",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1030, deposit_cents=480)],
            deposit_returns=[DepositReturnCreate(name="Leergut Kiste 4,80", quantity=1, deposit_cents=480)],
        ),
    )

    assert session.query(OpenItem).one().amount_cents == 1030
    sheet = load_workbook(document.excel_path, data_only=True).active
    assert sheet["F40"].value == -4.8
    assert sheet["F43"].value == 10.3


def test_create_document_uses_latest_customer_excel_as_working_template(session, tmp_path: Path):
    customer_folder = tmp_path / "Kunden" / "Metzgerei Karl"
    customer_folder.mkdir(parents=True)
    old_invoice = customer_folder / "_ RE 0525 Metzgerei Karl .xlsx"
    old_invoice.write_bytes((INPUT_DIR / "_ RE 0525 Metzgerei Karl .xlsx").read_bytes())
    customer = create_customer(
        session,
        CustomerCreate(name="Metzgerei Karl", folder_path=str(customer_folder), payment_method="SEPA"),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RE-FORT",
            delivery_date="2026-07-05",
            line_items=[
                DocumentLineItem(
                    name="Frucade Colamix 20x0,5",
                    quantity=4,
                    unit_price_cents=1048,
                    deposit_cents=310,
                )
            ],
            deposit_returns=[DepositReturnCreate(name="Pfand 3,10 EUR", quantity=3, deposit_cents=310)],
        ),
    )

    sheet = load_workbook(document.excel_path, data_only=False).active
    assert sheet["B13"].value == "Labertaler Apfelschorle 20x0,5"
    assert sheet["A13"].value is None
    assert sheet["B14"].value == "Frucade Colamix 20x0,5"
    assert sheet["A14"].value == 4
    assert sheet["B28"].value == "Frucade Zitrone 20x0,5"
    assert sheet["A28"].value is None
    assert sheet["A34"].value == 3
    assert sheet["C34"].value == -3.1


def test_create_invoice_rejects_duplicate_invoice_number(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(
            name="Cafe Nord",
            folder_path=str(tmp_path / "Kunden" / "Cafe Nord"),
        ),
    )
    payload = DocumentCreate(
        customer_id=customer.id,
        document_type="Rechnung",
        document_number="RG-1003",
        delivery_date="2026-06-21",
        line_items=[
            DocumentLineItem(
                name="Wasser 0,7",
                quantity=1,
                unit_price_cents=1299,
                deposit_cents=330,
            )
        ],
    )
    create_document(session, payload)

    try:
        create_document(session, payload)
    except ValueError as error:
        assert "Rechnungsnummer ist bereits vorhanden" in str(error)
    else:
        raise AssertionError("Duplicate invoice number was accepted")


def test_latest_invoice_number_returns_most_recent_invoice(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Cafe Nord", folder_path=str(tmp_path / "Kunden" / "Cafe Nord")),
    )
    for number in ("RG-1004", "RG-1005"):
        create_document(
            session,
            DocumentCreate(
                customer_id=customer.id,
                document_type="Rechnung",
                document_number=number,
                delivery_date="2026-06-21",
                line_items=[
                    DocumentLineItem(
                        name="Wasser 0,7",
                        quantity=1,
                        unit_price_cents=1299,
                        deposit_cents=330,
                    )
                ],
            ),
        )

    assert latest_invoice_number(session) == "RG-1005"


def test_create_invoice_copies_pdf_to_datev_upload_folder(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Cafe Nord", folder_path=str(tmp_path / "Kunden" / "Cafe Nord")),
    )
    datev_dir = tmp_path / "DATEV"

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-1006",
            delivery_date="2026-06-21",
            line_items=[
                DocumentLineItem(
                    name="Wasser 0,7",
                    quantity=1,
                    unit_price_cents=1299,
                    deposit_cents=330,
                )
            ],
        ),
        datev_upload_dir=datev_dir,
    )

    assert document.datev_export_path == str(datev_dir / "2026-06" / "2026-06-21_RE_RG-1006_Cafe_Nord.pdf")
    assert Path(document.datev_export_path).exists()


def test_create_document_can_generate_only_excel_first(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Excel Kunde", folder_path=str(tmp_path / "Kunden" / "Excel Kunde")),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-2001",
            delivery_date="2026-06-21",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299, deposit_cents=330)],
        ),
        assets={"excel"},
    )

    assert Path(document.excel_path).exists()
    assert not Path(document.pdf_path).exists()
    assert session.query(OpenItem).one().document_number == "RG-2001"


def test_create_document_can_generate_pdf_later_without_duplicate_invoice(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="PDF Kunde", folder_path=str(tmp_path / "Kunden" / "PDF Kunde")),
    )
    payload = DocumentCreate(
        customer_id=customer.id,
        document_type="Rechnung",
        document_number="RG-2002",
        delivery_date="2026-06-21",
        line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1299, deposit_cents=330)],
    )

    excel_document = create_document(session, payload, assets={"excel"})
    pdf_document = create_document(session, payload, assets={"pdf"})

    assert pdf_document.id == excel_document.id
    assert Path(pdf_document.excel_path).exists()
    assert Path(pdf_document.pdf_path).exists()
    assert session.query(OpenItem).count() == 1


def test_invoice_for_bank_transfer_customer_gets_due_footer_and_open_item_due_date(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(
            name="Ueberweiser Kunde",
            folder_path=str(tmp_path / "Kunden" / "Ueberweiser Kunde"),
            payment_method="Ueberweisung",
        ),
    )

    document = create_document(
        session,
        DocumentCreate(
            customer_id=customer.id,
            document_type="Rechnung",
            document_number="RG-U-1",
            delivery_date="2026-06-24",
            line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1000, deposit_cents=0)],
        ),
    )

    open_item = session.query(OpenItem).one()
    assert document.footer_text == "Bitte ueberweisen Sie den Rechnungsbetrag bis zum 2026-07-01."
    assert open_item.document_date == "2026-06-24"
    assert open_item.due_date == "2026-07-01"
    assert open_item.payment_method == "Ueberweisung"


def test_partial_document_generation_rejects_changed_payload_after_first_asset(session, tmp_path: Path):
    customer = create_customer(
        session,
        CustomerCreate(name="Teil Export Kunde", folder_path=str(tmp_path / "Kunden" / "Teil Export Kunde")),
    )
    first_payload = DocumentCreate(
        customer_id=customer.id,
        document_type="Rechnung",
        document_number="RG-PART-1",
        delivery_date="2026-06-24",
        line_items=[DocumentLineItem(name="Wasser", quantity=1, unit_price_cents=1000, deposit_cents=0)],
    )
    changed_payload = DocumentCreate(
        customer_id=customer.id,
        document_type="Rechnung",
        document_number="RG-PART-1",
        delivery_date="2026-06-24",
        line_items=[DocumentLineItem(name="Wasser", quantity=2, unit_price_cents=1000, deposit_cents=0)],
    )

    create_document(session, first_payload, assets={"excel"})

    try:
        create_document(session, changed_payload, assets={"pdf"})
    except ValueError as error:
        assert "Belegdaten passen nicht" in str(error)
    else:
        raise AssertionError("Changed partial document payload was accepted")
